import argparse
import csv
import io
import json
import re
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import extract_msg
from openpyxl import load_workbook


def clean_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_header(value):
    return re.sub(r"[^a-z0-9]+", "_", clean_text(value).lower()).strip("_")


def cell_value(value):
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


# The ERP "Screen" export is a legacy fixed-width report: a few decorative
# divider/title rows before the real header, which we find by looking for
# the one column every reconciliation needs -- PAV (the shared reference
# with FloraHolland's invoices).
def parse_erp(input_path: Path):
    workbook = load_workbook(filename=str(input_path), data_only=True, read_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    all_rows = [list(row) for row in worksheet.iter_rows(values_only=True)]

    header_row_index = None
    headers = []
    for idx, row in enumerate(all_rows):
        normalized = [normalize_header(value) for value in row]
        if "pav" in normalized:
            header_row_index = idx
            headers = normalized
            break

    if header_row_index is None:
        return {"rows": [], "error": "Could not find a header row containing PAV in this file"}

    rows = []
    for row in all_rows[header_row_index + 1:]:
        if not any(clean_text(value) for value in row):
            continue
        record = {}
        for col_index, header in enumerate(headers):
            if not header or col_index >= len(row):
                continue
            record[header] = cell_value(row[col_index])
        rows.append(record)

    return {"rows": rows}


TYPE_CODE_RE = re.compile(r"[._](FK|FC|HA)[._]", re.IGNORECASE)


# The XML-folder messages ("XCrossIndustryInvoice;049876.FC.2026.0181") don't
# spell out "Klokfactuur"/"Connect factuur"/"Handel aankopen" in the subject
# the way the PDF-folder messages do -- but both always carry the same
# FK/FC/HA type code from the invoice number, so that's the reliable signal.
def classify_subject(subject: str) -> str:
    normalized = (subject or "").lower()
    match = TYPE_CODE_RE.search(subject or "")
    if match:
        code = match.group(1).upper()
        return {"FK": "klokfactuur", "FC": "connect", "HA": "handel"}[code]
    if "klokfactuur" in normalized:
        return "klokfactuur"
    if "connect factuur" in normalized:
        return "connect"
    if "handel aankopen" in normalized:
        return "handel"
    return "other"


INVOICE_NUMBER_RE = re.compile(r"(\d{6}[._][A-Z]{2}[._]\d{4}[._]\d{3,4})", re.IGNORECASE)


def extract_invoice_number(subject: str):
    match = INVOICE_NUMBER_RE.search(subject or "")
    if not match:
        return ""
    return match.group(1).replace("_", ".").upper()


def find_xml_bytes(msg):
    for attachment in msg.attachments:
        name = (attachment.longFilename or attachment.shortFilename or "").lower()
        if not attachment.data:
            continue
        if name.endswith(".xml"):
            return attachment.data
        if name.endswith(".zip"):
            try:
                with zipfile.ZipFile(io.BytesIO(attachment.data)) as archive:
                    xml_names = [n for n in archive.namelist() if n.lower().endswith(".xml")]
                    if xml_names:
                        return archive.read(xml_names[0])
            except zipfile.BadZipFile:
                continue
    return None


def local_tag(tag: str) -> str:
    return tag.split("}")[-1] if "}" in tag else tag


def find_child_text(element, tag_name):
    for child in element.iter():
        if local_tag(child.tag) == tag_name:
            return (child.text or "").strip()
    return None


def find_reference(element, scheme_name):
    for ref in element.iter():
        if local_tag(ref.tag) != "ReferenceReferencedDocument":
            continue
        for child in ref.iter():
            if local_tag(child.tag) == "IssuerAssignedID" and child.attrib.get("schemeName") == scheme_name:
                return (child.text or "").strip()
    return None


def parse_number(text):
    if text is None:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def find_supplier_info(item):
    for party in item.iter():
        if local_tag(party.tag) != "SupplierParty":
            continue
        gln = None
        fh_number = None
        name = None
        for child in party.iter():
            if local_tag(child.tag) == "PrimaryID" and gln is None:
                gln = (child.text or "").strip()
            elif local_tag(child.tag) == "AdditionalID" and child.attrib.get("schemeAgencyName") == "FH":
                fh_number = (child.text or "").strip()
            elif local_tag(child.tag) == "Name" and name is None:
                name = (child.text or "").strip()
        return {"gln": gln or "", "fh_number": fh_number or "", "name": name or ""}
    return {"gln": "", "fh_number": "", "name": ""}


# Header-level element (one per invoice, appears once before any line item)
# identifying which of our own buying entities this invoice was billed to --
# e.g. "Juniflor Flower Export BV". Distinct from SupplierParty, which is the
# grower/seller and is scoped per line item.
def find_invoicee_info(root):
    for party in root.iter():
        if local_tag(party.tag) != "InvoiceeParty":
            continue
        gln = None
        name = None
        for child in party.iter():
            if local_tag(child.tag) == "PrimaryID" and gln is None:
                gln = (child.text or "").strip()
            elif local_tag(child.tag) == "Name" and name is None:
                name = (child.text or "").strip()
        return {"gln": gln or "", "name": name or ""}
    return {"gln": "", "name": ""}


def parse_invoice_xml(xml_bytes):
    root = ET.fromstring(xml_bytes)
    invoicee = find_invoicee_info(root)
    # Header-level issue date, used as a fallback when a line has no date of
    # its own -- real dates from the XML instead of the date manually typed
    # into the upload form, which is what a calendar/day report needs to be
    # trustworthy.
    issue_date = find_child_text(root, "IssueDateTime") or ""
    lines = []
    for item in root.iter():
        if local_tag(item.tag) != "InvoiceTradeLineItem":
            continue
        quantity = parse_number(find_child_text(item, "BilledQuantity"))
        description = find_child_text(item, "DescriptionText") or ""
        unit_price = parse_number(find_child_text(item, "ChargeAmount"))
        total = parse_number(find_child_text(item, "GrandTotalAmount"))
        reference_bt = find_reference(item, "BT") or ""
        supplier = find_supplier_info(item)
        line_date = find_child_text(item, "LineDateTime") or issue_date
        lines.append({
            "description": description,
            "quantity": quantity,
            "unit_price": unit_price,
            "total": total,
            "reference_bt": reference_bt,
            "supplier_gln": supplier["gln"],
            "supplier_fh_number": supplier["fh_number"],
            "supplier_name": supplier["name"],
            "company_name": invoicee["name"],
            "line_date": line_date,
        })
    return lines


def normalize_supplier_name(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", (name or "").lower())


# Master data dumps ("stamgegevens") -- semicolon-delimited, ~240 columns,
# but only Code/Naam/GLN kweke matter here. Growers ("kwekers") and direct-
# trade partners ("leveranciers") are separate exports with different
# coverage, so both are read and merged (leveranciers entries win on
# conflict since they're the more specific, purchase-relevant list).
#
# GLN is the primary, exact join key, but not every row has one (some
# direct-trade partners are only recorded as an internal alias with a name,
# no GLN at all) -- for those, a normalized-name index is built too, used
# only as a *suggestion* on the "supplier not linked" screen, never to
# silently auto-match, since company names can collide or vary in ways a
# GLN can't.
def load_supplier_csv(path: Path, source_label: str):
    by_gln = {}
    by_name = {}
    with open(path, encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.reader(handle, delimiter=";")
        header = next(reader, [])
        idx = {name: i for i, name in enumerate(header)}
        code_idx = idx.get("Code")
        naam_idx = idx.get("Naam")
        gln_idx = idx.get("GLN kweke")
        if code_idx is None or naam_idx is None or gln_idx is None:
            return by_gln, by_name
        for row in reader:
            if len(row) <= max(code_idx, naam_idx, gln_idx):
                continue
            code = row[code_idx].strip()
            name = row[naam_idx].strip()
            if not code:
                continue
            gln = row[gln_idx].strip()
            if gln:
                by_gln[gln] = {"code": code, "name": name, "source": source_label}
            normalized_name = normalize_supplier_name(name)
            if normalized_name:
                by_name[normalized_name] = {"code": code, "name": name, "source": source_label}
    return by_gln, by_name


def parse_suppliers(kwekers_path: Path, leveranciers_path: Path):
    by_gln = {}
    by_name = {}
    for path, label in [(kwekers_path, "kwekers"), (leveranciers_path, "leveranciers")]:
        if not path:
            continue
        gln_entries, name_entries = load_supplier_csv(path, label)
        by_gln.update(gln_entries)
        by_name.update(name_entries)
    return {"by_gln": by_gln, "by_name": by_name}


# Only Klokfactuur/Connect/Handel messages carry the CII XML this tool needs
# -- anything else (e.g. an AI2 Dagnota) is reported as skipped so the
# caller can show why a file in the upload wasn't used, rather than it
# silently disappearing.
def parse_veiling(input_path: Path):
    extract_dir = Path(tempfile.mkdtemp(prefix="inkoop-veiling-"))
    with zipfile.ZipFile(input_path) as archive:
        archive.extractall(extract_dir)

    invoices = []
    skipped = []
    for msg_path in sorted(extract_dir.rglob("*.msg")):
        msg = extract_msg.Message(str(msg_path))
        try:
            subject = msg.subject or ""
            kind = classify_subject(subject)
            if kind == "other":
                skipped.append({"file_name": msg_path.name, "reason": f"Unsupported message type (subject: {subject})"})
                continue

            xml_bytes = find_xml_bytes(msg)
            if not xml_bytes:
                skipped.append({"file_name": msg_path.name, "reason": "No invoice XML attachment found"})
                continue

            lines = parse_invoice_xml(xml_bytes)
            invoice_number = extract_invoice_number(subject)
            invoices.append({
                "file_name": msg_path.name,
                "type": kind,
                "invoice_number": invoice_number,
                "supplier_number": invoice_number.split(".")[0] if invoice_number else "",
                "subject": subject,
                "lines": lines,
            })
        except Exception as error:  # noqa: BLE001 -- report per-file, never abort the whole batch
            skipped.append({"file_name": msg_path.name, "reason": str(error)})
        finally:
            msg.close()

    return {"invoices": invoices, "skipped": skipped}


def main():
    parser = argparse.ArgumentParser()
    subparsers = parser.add_subparsers(dest="command", required=True)

    erp_parser = subparsers.add_parser("parse-erp")
    erp_parser.add_argument("--input", required=True)

    veiling_parser = subparsers.add_parser("parse-veiling")
    veiling_parser.add_argument("--input", required=True)

    suppliers_parser = subparsers.add_parser("parse-suppliers")
    suppliers_parser.add_argument("--kwekers", required=False, default="")
    suppliers_parser.add_argument("--leveranciers", required=False, default="")

    args = parser.parse_args()

    if args.command == "parse-erp":
        print(json.dumps(parse_erp(Path(args.input))))
    elif args.command == "parse-veiling":
        print(json.dumps(parse_veiling(Path(args.input))))
    elif args.command == "parse-suppliers":
        kwekers_path = Path(args.kwekers) if args.kwekers else None
        leveranciers_path = Path(args.leveranciers) if args.leveranciers else None
        print(json.dumps(parse_suppliers(kwekers_path, leveranciers_path)))


if __name__ == "__main__":
    main()
