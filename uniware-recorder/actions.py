"""Named Uniware actions: record once, replay from the dashboard.

An action is a list of steps, each one key (combination) or typed text plus
the screen Uniware should show afterwards (its Prg / [number] id, see
screen_reader.screen_identity). Replay sends the same keys to the Uniware
window and checks after every step that the expected screen is there --
anything else stops the run, so it never keeps typing into a wrong screen.

Steps can capture the screen (its text, or every page via Page Down) into
outputs/, and typed text can be turned into a field ("{invoice}") whose
value is asked when the action is started.
"""
from __future__ import annotations

import csv
import ctypes
import asyncio
import difflib
import datetime
import json
import os
import queue
import re
import subprocess
import threading
import time
from ctypes import wintypes
from pathlib import Path

from pynput import keyboard

import recorder
import screen_reader
from PIL import Image, ImageChops

APP_DIR = Path(__file__).resolve().parent
ACTIONS_DIR = APP_DIR / "actions"
OUTPUTS_DIR = APP_DIR / "outputs"
DUMP_PARTS_DIR = OUTPUTS_DIR / "dump-parts"
# Uniware's "loading" box: "Gegevens verzamelen ... / Moment geduld aub..."
BUSY_RE = re.compile(r"moment\s*geduld|gegevens\s*verzamelen", re.IGNORECASE)


def _relative_sheet_paths(path: Path):
    """openpyxl writes the workbook's sheet links as "/xl/worksheets/..."; the
    files from Uniware (and what SnappySjaak's reader expects) have them
    relative: "worksheets/...". Rewrite them that way."""
    import zipfile
    name = "xl/_rels/workbook.xml.rels"
    with zipfile.ZipFile(path) as source:
        entries = [(info, source.read(info.filename)) for info in source.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for info, data in entries:
            if info.filename == name:
                data = data.replace(b'Target="/xl/', b'Target="')
            target.writestr(info, data)


def _trimmed(row) -> list:
    """A spreadsheet row without its empty cells at the end."""
    values = list(row)
    while values and values[-1] in (None, ""):
        values.pop()
    return values
user32 = ctypes.windll.user32

# The Uniware window: "Connect (C Uniware Computer Systems BV) (Session 1: ...)".
DEFAULT_WINDOW_TITLE = "Uniware Computer Systems"
TYPING_GAP_SECONDS = 0.7  # typed characters closer together than this are one step
PAGE_WAIT_SECONDS = 8  # Page Down: no change within this time = end of the list


def default_max_wait(recorded_seconds: float) -> int:
    """How long a run waits for a step's screen: plenty more than it took
    while recording (Uniware can be slow), at least 30 s."""
    return int(min(600, max(30, recorded_seconds * 4)))


def slugify(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")[:60] or "action"


def list_actions() -> list[dict]:
    ACTIONS_DIR.mkdir(parents=True, exist_ok=True)
    actions = []
    for path in sorted(ACTIONS_DIR.glob("*.json")):
        try:
            actions.append(json.loads(path.read_text(encoding="utf-8")))
        except (OSError, json.JSONDecodeError):
            continue
    return actions


def load_action(slug: str) -> dict | None:
    path = ACTIONS_DIR / f"{slug}.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else None


def save_action(action: dict):
    ACTIONS_DIR.mkdir(parents=True, exist_ok=True)
    path = ACTIONS_DIR / f"{action['slug']}.json"
    temp = path.with_suffix(".tmp")
    temp.write_text(json.dumps(action, indent=2, ensure_ascii=False), encoding="utf-8")
    temp.replace(path)


DATE_WORDS = ("today", "tomorrow", "yesterday", "next_workday", "last_workday")


def date_for_word(word: str, now: datetime.date | None = None) -> datetime.date | None:
    """today / tomorrow / yesterday / next_workday / last_workday (the
    next / previous Monday-Friday) -> that date; anything else -> None."""
    word = str(word or "").strip().lower().replace(" ", "_")
    today = now or datetime.date.today()
    if word == "today":
        return today
    if word == "tomorrow":
        return today + datetime.timedelta(days=1)
    if word == "yesterday":
        return today - datetime.timedelta(days=1)
    if word in ("next_workday", "last_workday"):
        step = 1 if word == "next_workday" else -1
        day = today + datetime.timedelta(days=step)
        while day.weekday() >= 5:
            day += datetime.timedelta(days=step)
        return day
    return None


def parse_date_value(value: str, now: datetime.date | None = None) -> datetime.date | None:
    """A date as someone types it when starting an action: a word (see
    date_for_word) or 14-10-2026 / 14-10-26 / 2026-10-14 / 14/10/2026 /
    14102026 / 141026."""
    text = str(value or "").strip()
    day = date_for_word(text, now)
    if day:
        return day
    patterns = [
        (r"(\d{4})-(\d{1,2})-(\d{1,2})", ("y", "m", "d")),
        (r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})", ("d", "m", "y")),
        (r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{2})", ("d", "m", "y2")),
        (r"(\d{2})(\d{2})(\d{4})", ("d", "m", "y")),
        (r"(\d{2})(\d{2})(\d{2})", ("d", "m", "y2")),
    ]
    for pattern, order in patterns:
        found = re.fullmatch(pattern, text)
        if not found:
            continue
        parts = dict(zip(order, (int(group) for group in found.groups())))
        year = parts.get("y") or 2000 + parts["y2"]
        try:
            return datetime.date(year, parts["m"], parts["d"])
        except ValueError:
            return None
    return None


def format_date_like(day: datetime.date, recorded: str) -> str:
    """The date written like the one typed while recording ("20-10-2026",
    "20102026", "201026", "20-10-26 13:00"...): same separator and year
    length, and the rest of the recorded text kept. Without a recorded date:
    dd-mm-yyyy."""
    # With separators first ("20-10-202613:00:00": the time may follow the
    # year straight away), then digits only as a whole ("20102026").
    found = (re.search(r"(?<!\d)(\d{2})([-/.])(\d{2})\2(\d{4}|\d{2}(?!\d))", recorded or "")
             or re.search(r"(?<!\d)(\d{2})()(\d{2})(\d{4}|\d{2})(?!\d)", recorded or ""))
    if not found:
        return day.strftime("%d-%m-%Y")
    separator, year = found.group(2), found.group(4)
    text = f"{day:%d}{separator}{day:%m}{separator}{day:%Y}" if len(year) == 4 else f"{day:%d}{separator}{day:%m}{separator}{day:%y}"
    return recorded[:found.start()] + text + recorded[found.end():]


def action_fields(action: dict) -> list[str]:
    fields = []
    # Filled per order by a "for each order" step -- not asked at start.
    filled = {step["foreach"].get("field") or "order" for step in action.get("steps", []) if step.get("foreach")}
    for step in action.get("steps", []):
        for name in re.findall(r"\{([a-z0-9_]+)\}", (step.get("text") or "") + (step.get("pick") or "")):
            if name not in fields and name not in filled and name not in DATE_WORDS:
                fields.append(name)
    return fields


# ------------------------------------------------------------------ menus
# Uniware's menus are levels side by side: the main menu on the left, each
# Right arrow opens the blue (selected) item's submenu as the next column,
# Left goes back a level. The level you're in is the rightmost column with a
# blue item (a column further right without one is just a preview of the
# selected item's submenu). So a menu route is the blue item per level --
# ["Expeditie", "Inklaring bedrijf"] -- and replay steers there by reading
# the screen, whatever the menu looked like when it started.
MENU_KEYS = {"Down", "Up", "Left", "Right", "Page Down", "Page Up", "Home", "End"}


def menu_levels(layout: list) -> list[dict]:
    """The columns of a menu screen, left to right: {x, items, selected}
    (selected = index of the blue item, or None)."""
    levels = []
    for line in sorted(layout, key=lambda line: line["x"]):
        if levels and line["x"] - levels[-1]["x"] <= 40:
            levels[-1]["lines"].append(line)
        else:
            levels.append({"x": line["x"], "lines": [line]})
    result = []
    for level in levels:
        lines = sorted(level["lines"], key=lambda line: line["y"])
        selected = [index for index, line in enumerate(lines) if line.get("highlighted")]
        result.append({"x": level["x"], "items": [line["text"] for line in lines], "selected": selected[0] if selected else None})
    return result


# The main menu's items: a real menu screen always shows them as its first
# column (a program screen such as Prg. 423 Transport afspraken has a Prg
# number and a long first column too, and was taken for a menu).
MAIN_MENU_ITEMS = ["Offertes", "Bestellingen", "Inkoop", "Verkoop", "Voorraad", "Productie", "Expeditie", "Fust administratie",
                   "Communicatie", "Onderhoud gegevens", "Afdrukken gegevens", "Financieel", "Systeem onderhoud"]


def is_menu(screen_id: str, layout: list) -> bool:
    """A menu screen: a Prg number and the main menu as the first column."""
    levels = menu_levels(layout)
    if not (screen_id.startswith("prg-") and levels and len(levels[0]["items"]) >= 8 and levels[0]["selected"] is not None):
        return False
    known = sum(1 for item in levels[0]["items"] if any(same_item(item, name) for name in MAIN_MENU_ITEMS))
    return known >= 6


def menu_path(layout: list) -> list[str]:
    """The blue item of every level up to the one you're in."""
    levels = menu_levels(layout)
    active = max((index for index, level in enumerate(levels) if level["selected"] is not None), default=-1)
    return [levels[index]["items"][levels[index]["selected"]] if levels[index]["selected"] is not None else "" for index in range(active + 1)]


def _screen_number(screen_id: str):
    match = re.fullmatch(r"(prg|scr)-(\d+)", screen_id or "")
    return (match.group(1), match.group(2)) if match else None


def same_screen(expected_id: str, expected_title: str, seen_id: str, seen_title: str) -> bool:
    """Is Uniware on the expected screen? The screen name comes from text
    recognition, which can misread a letter or the closing bracket
    ("Customsclearance cornpany [86]" / "company [861]"): same number, or
    one digit more at the end, plus a nearly identical title is the same
    screen. Different numbers otherwise never are (Prg. 000 vs 420)."""
    if expected_id == seen_id:
        return True
    letters = lambda text: re.sub(r"[^a-z]+", "", re.sub(r"\[[^\]]*\]?|prg\.?\s*\S+", "", (text or "").lower()))
    similar = difflib.SequenceMatcher(None, letters(expected_title), letters(seen_title)).ratio() >= 0.85
    expected_number, seen_number = _screen_number(expected_id), _screen_number(seen_id)
    if expected_number and seen_number:
        if expected_number[0] != seen_number[0]:
            return False
        short, long = sorted((expected_number[1], seen_number[1]), key=len)
        return similar and long.startswith(short) and len(long) - len(short) <= 1
    if similar and letters(expected_title):
        return True
    # One of the two names was read badly while recording, e.g. "o |
    # Customsclearance lines | Info | Filter Settings | ..." (the whole top
    # of the screen) for "Customsclearance lines [85]": the shorter title
    # inside the longer one is the same screen.
    short, long = sorted((letters(expected_title), letters(seen_title)), key=len)
    return len(short) >= 10 and short in long


def _plain(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", text.lower())


def same_item(seen: str, wanted: str) -> bool:
    """Menu text as read now vs as recorded -- OCR may differ a little
    ("vt_control" / "vt control", "virtuee" / "virtueel")."""
    left, right = _plain(seen), _plain(wanted)
    if not left or not right:
        return False
    if left == right or (min(len(left), len(right)) >= 5 and (left.startswith(right) or right.startswith(left))):
        return True
    return difflib.SequenceMatcher(None, left, right).ratio() >= 0.85


# ------------------------------------------------------------------ day lists
# A day list carries information from one action to the next (e.g. the
# accessoire orders of tour 515: collected by one action, checked for SPICER
# by the next, assigned to tour 515 / 516 by the last). One JSON file per list
# in daylists/; each item is keyed by its order number.
DAYLISTS_DIR = APP_DIR / "daylists"


def day_list_slug(name: str) -> str:
    return slugify(name) or "list"


def load_day_list(name: str) -> dict:
    path = DAYLISTS_DIR / f"{day_list_slug(name)}.json"
    if path.exists():
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            pass
    return {"name": name, "updated": "", "items": []}


def save_day_list(data: dict):
    DAYLISTS_DIR.mkdir(parents=True, exist_ok=True)
    data["updated"] = datetime.datetime.now().isoformat(timespec="seconds")
    path = DAYLISTS_DIR / f"{day_list_slug(data['name'])}.json"
    temp = path.with_suffix(".tmp")
    temp.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    temp.replace(path)


def list_day_lists() -> list[dict]:
    DAYLISTS_DIR.mkdir(parents=True, exist_ok=True)
    lists = []
    for path in sorted(DAYLISTS_DIR.glob("*.json")):
        try:
            lists.append(json.loads(path.read_text(encoding="utf-8")))
        except (OSError, json.JSONDecodeError):
            continue
    return lists


def day_list_item(data: dict, key: str) -> dict | None:
    return next((item for item in data.get("items", []) if item.get("key") == key), None)


def row_order_number(text: str) -> str | None:
    """The order number of a list row (6-8 digits, look-alike letters read as
    digits): 4840711 in "508 EC1211 4840711 0.00 0.05"."""
    text = text.translate(_DIGIT_LOOKALIKES)
    numbers = re.findall(r"(?<![\w.,])\d{6,8}(?![\w.,])", text)
    if numbers:
        return numbers[0]
    # OCR sometimes glues the customer code to it: "EF3174841449" -> the
    # last 7 digits of the first long digit run.
    glued = re.findall(r"(?<![.,])\d{9,12}(?![\w.,])", text)
    return glued[0][-7:] if glued else None


def list_row_order(text: str) -> str | None:
    """The order number of a list row that starts with a customer code:
    "EF317 4841449 0.00 0.11 ..." / "508 EC1211 4840711 0.00" -> the number
    after the code, before the first amount (None for any other line)."""
    found = re.search(r"(?<!\S)([A-Za-z][A-Za-z0-9]{2,7})\s+(\S{7})(?:\s+\S*\d[.,]\d\d|\s*$)", text)
    if not found:
        return None
    number = found.group(2).translate(_DIGIT_LOOKALIKES)
    return number if number.isdigit() else None


def list_row_lot(text: str) -> str | None:
    """The Connect lot in a list row: an 8-digit number starting with 4
    ("... 44759254 J 2271237 ..."), look-alike letters read as digits."""
    for token in str(text or "").split():
        digits = token.translate(_DIGIT_LOOKALIKES)
        if re.fullmatch(r"4\d{7}", digits):
            return digits
    return None


def list_row_brief(text: str) -> str | None:
    """The brief number in a list row's Info column ("... 5 FA3CWD/A 44858309
    J 10", "X01907/a", "407971/-") -- a row without one (Info empty, or not
    readable) is a lot that can't be matched yet."""
    found = re.search(r"(?<!\S)([A-Za-z0-9]{5,8}/[A-Za-z~_-]?)(?!\S)", str(text or ""))
    if not found or not re.search(r"\d", found.group(1)):
        return None
    return found.group(1)


def popup_box(before, after):
    """Where a popup opened: the biggest block of rows that changed between
    the screenshot before and after its key (x0, y0, x1, y1), or None. Only
    its lines count -- not the list around it."""
    if before is None or after is None or before.size != after.size:
        return None
    mask = ImageChops.difference(before.convert("RGB"), after.convert("RGB")).convert("L").point(lambda v: 255 if v > 40 else 0)
    if not mask.getbbox():
        return None
    width, height = mask.size
    # Mean per pixel row: text rows of a popup change a few % of the width;
    # the blank space between its lines is bridged (up to 18 px).
    rows = list(mask.resize((1, height), Image.BOX).getdata())
    best, start, end, gap = (0, -1), None, None, 0
    for y, value in enumerate(rows + [0] * 20):
        if value > 3:
            start = y if start is None else start
            end, gap = y, 0
        elif start is not None:
            gap += 1
            if gap > 18:
                if end - start > best[1] - best[0]:
                    best = (start, end)
                start = None
    top, bottom = best
    if bottom - top < 40:
        return None
    columns = list(mask.crop((0, top, width, bottom + 1)).resize((width, 1), Image.BOX).getdata())
    xs = [x for x, value in enumerate(columns) if value > 3]
    return (xs[0], top, xs[-1], bottom) if xs else None


def screens_differ(before, after) -> bool:
    """More changed between two screenshots than a blinking caret."""
    if before is None or after is None or before.size != after.size:
        return True
    box = ImageChops.difference(before.convert("RGB"), after.convert("RGB")).convert("L").point(lambda v: 255 if v > 40 else 0).getbbox()
    return bool(box) and (box[2] - box[0]) * (box[3] - box[1]) > 600


def read_cursor_row_key(image, entry, row_key):
    """The teal cursor row read again, enlarged and in grey (OCR misreads
    digits on teal): its key, or None."""
    if image is None or entry is None:
        return None
    top = max(0, int(entry["y"]) - 6)
    crop = image.crop((0, top, image.width, min(image.height, int(entry["y"]) + 24))).convert("L")
    crop = crop.resize((crop.width * 2, crop.height * 2), Image.LANCZOS).convert("RGB")
    try:
        words = asyncio.run(screen_reader._ocr(crop))
    except Exception:  # noqa: BLE001 -- a fallback only
        return None
    return row_key(" ".join(word["text"] for word in words))


def read_screen_value(label: str, lines: list[str]) -> str | None:
    """The number after a label on the screen: "EPP: [9100185445726]" ->
    "9100185445726" (look-alike letters read as digits)."""
    if not label:
        return None
    found = re.search(re.escape(label) + r"\s*[:=]?\s*\[?\s*([0-9OoSsIlBDZ]{6,16})", "\n".join(lines), re.IGNORECASE)
    if not found:
        return None
    digits = found.group(1).translate(_DIGIT_LOOKALIKES)
    return digits if digits.isdigit() else None


def normalize_foreach(settings: dict) -> dict:
    try:
        count = max(1, min(50, int(float(settings.get("count") or 1))))
    except (TypeError, ValueError):
        count = 1
    field = re.sub(r"[^a-z0-9_]+", "_", str(settings.get("field") or "order").lower()).strip("_") or "order"
    return {
        "list": str(settings.get("list") or "").strip(),
        "field": field,
        "count": count,
        # Only items that have no label yet (a run can be repeated / resumed).
        "only_open": settings.get("only_open", True) is not False,
    }


def normalize_classify(settings: dict) -> dict:
    return {
        "list": str(settings.get("list") or "").strip(),
        "word": str(settings.get("word") or "").strip(),
        "label_if": str(settings.get("label_if") or "").strip(),
        "label_else": str(settings.get("label_else") or "").strip(),
    }


# A product line on an order screen starts with its line number and article
# number ("1 44752669a ...") or has a quantity like "1 x 50" / "2 x1" (OCR
# reads x1 as "xl").
PRODUCT_LINE_RE = re.compile(r"^\s*\d{1,3}\s+\d{6,9}[a-z]?\s|\d+\s*x\s*[\dlI|]", re.IGNORECASE)


# ------------------------------------------------------------------ list loop
# A list loop ends after this many presses of the next key in a row that
# don't bring a new row (one can be a slow screen).
LIST_LOOP_END_TRIES = 3


def list_loop_keys(value: str) -> list[str]:
    """ "Delete, Right, Enter" -> ["Delete", "Right", "Enter"] (a comma inside
    a combo like "Ctrl+," is not supported)."""
    return [part.strip() for part in str(value or "").split(",") if part.strip()]


def normalize_list_loop(loop: dict) -> dict:
    """Settings of a list loop step (see ActionRun._run_list_loop)."""
    words = lambda value: [word.strip() for word in (value if isinstance(value, list) else str(value or "").split(",")) if str(word).strip()]
    key = lambda value, default: str(value or "").strip() or default
    for name, default in (("match_key", "Delete"), ("next_key", "Down")):
        # Several keys in a row, comma-separated ("Delete, Right, Enter":
        # Delete, then answer Uniware's confirmation) -- each must be a key.
        for part in list_loop_keys(key(loop.get(name), default)):
            try:
                parse_combo(part)
            except (AttributeError, KeyError, ValueError):
                raise RuntimeError(f"List loop: '{part}' is not a key (e.g. Delete, Right, Enter, Down).")
    try:
        max_rows = max(1, min(2000, int(float(loop.get("max_rows") or 500))))
    except (TypeError, ValueError):
        max_rows = 500
    return {
        "match": words(loop.get("match")),
        "keep": words(loop.get("keep")),
        "match_key": key(loop.get("match_key"), "Delete"),
        "next_key": key(loop.get("next_key"), "Down"),
        "max_rows": max_rows,
        "dry_run": bool(loop.get("dry_run")),
        # The key doesn't take the row out of the list (e.g. Delete only
        # clears its tour number): move on with the next key afterwards.
        "row_stays": bool(loop.get("row_stays")),
        # Only rows whose text contains this (e.g. "508") are looked at --
        # a row already cleared is passed by (safe to run again).
        "only_rows_with": str(loop.get("only_rows_with") or "").strip(),
        # Collect mode: no keys but the next key -- every row's order number
        # (with the type word seen) goes into this day list (emptied first).
        "collect_into": str(loop.get("collect_into") or "").strip(),
        # What a collected row is known by: its order number (the 7 digits
        # after the customer code) or its Connect lot (8 digits, 44...).
        "collect_key": "lot" if loop.get("collect_key") == "lot" else "order",
        # Also read per row from the screen, e.g. "EPP" -> the number after
        # "EPP:" in the bottom bar for the cursor row; saved as a CSV too.
        "collect_value": str(loop.get("collect_value") or "").strip(),
        # Per row, before the next key: open a popup (e.g. "Shift+?" -- the
        # supplier info), keep every line it shows, close it (e.g. "Esc").
        "row_popup_keys": str(loop.get("row_popup_keys") or "").strip(),
        "row_close_keys": str(loop.get("row_close_keys") or "").strip() or "Esc",
        # Act on a day list instead of words on screen: a row matches when its
        # order number has this label in that list (e.g. 516).
        "use_list": str(loop.get("use_list") or "").strip(),
        "use_label": str(loop.get("use_label") or "").strip(),
    }


# ------------------------------------------------------------------ lists
ARROW_KEYS = {"Down", "Up", "Page Down", "Page Up"}


def screen_rows(layout: list) -> list[dict]:
    """Text lines grouped into screen rows (top to bottom): {text, y,
    highlighted}. OCR splits a list row into one line per column."""
    rows = []
    for line in sorted(layout, key=lambda line: line["y"]):
        if rows and abs(line["y"] - rows[-1]["y"]) <= 8:
            rows[-1]["parts"].append(line)
            rows[-1]["highlighted"] = rows[-1]["highlighted"] or line.get("highlighted", False)
            rows[-1]["cursor"] = rows[-1]["cursor"] or line.get("cursor", False)
        else:
            rows.append({"y": line["y"], "parts": [line], "highlighted": line.get("highlighted", False), "cursor": line.get("cursor", False)})
    # The selected row of a list is the teal cursor row -- the one teal
    # across the most fields (a page counter "(1) 1-15 of 15" is teal too);
    # only a screen without one falls back to blue text (blue is also used
    # for column headers).
    for row in rows:
        row["cursor_parts"] = sum(1 for part in row["parts"] if part.get("cursor"))
    most = max((row["cursor_parts"] for row in rows), default=0)
    return [{"y": row["y"], "highlighted": (row["cursor_parts"] == most) if most else row["highlighted"],
             "text": " ".join(part["text"] for part in sorted(row["parts"], key=lambda part: part["x"]))} for row in rows]


# Windows' text recognition reads Uniware's digits as letters when a number
# stands on its own: "508" -> "soa", "515" -> "SIS", "55" -> "SS". When
# looking for a number, read those letters as the digits they resemble.
_DIGIT_LOOKALIKES = str.maketrans({"S": "5", "s": "5", "O": "0", "o": "0", "Û": "0", "Ø": "0", "D": "0", "a": "8", "B": "8",
                                   "I": "1", "l": "1", "i": "1", "|": "1", "Z": "2", "z": "2", "G": "6", "b": "6", "q": "9", "g": "9"})


def _as_digits(text: str, target: str) -> str:
    return text.translate(_DIGIT_LOOKALIKES) if target.strip().isdigit() else text


def value_matches(text: str, target: str) -> bool:
    """Does `text` show `target` as a whole value ("508" matches "508 GBR"
    and an OCR'd "soa", not "1508" or "50,80")?"""
    pattern = re.compile(rf"(?<![\w.,]){re.escape(target.strip())}(?![\w.,]|,\d|\.\d)", re.IGNORECASE)
    return bool(pattern.search(text) or pattern.search(_as_digits(text, target)))


def find_row(rows: list[dict], target: str) -> int | None:
    """Index of the row showing `target` as a whole value; a row that starts
    with it wins."""
    matches = [index for index, row in enumerate(rows) if value_matches(row["text"], target)]
    starting = [index for index in matches if _as_digits(rows[index]["text"].strip(), target).lower().startswith(target.strip().lower())]
    return (starting or matches or [None])[0]


def _words_with_x(line: dict) -> list[tuple[float, float, str]]:
    """(left, right, word) of every word in a line -- by character position,
    the font is fixed-width."""
    width = line.get("w") or len(line["text"]) * 14
    char = width / max(1, len(line["text"]))
    return [(line["x"] + match.start() * char, line["x"] + match.end() * char, match.group())
            for match in re.finditer(r"\S+", line["text"])]


def read_column(layout: list, header: str, row_with: str = "") -> list[str]:
    """The values under a table column, e.g. "invoiceRes", one per row
    (empty cells left out). The column runs from the end of the header
    before it to the start of the header after it (headers can be narrower
    than their column, right-aligned like "order"). A header cut to the
    column width ("invoiceRe") is found too. row_with: only rows whose text
    contains that (e.g. "LTD" -- one company of several)."""
    wanted = header.strip().lower()
    best = None
    for line in layout:
        for left, right, word in _words_with_x(line):
            token = word.lower()
            if token == wanted or (len(token) >= 5 and wanted.startswith(token)):
                best = (line["y"], left, right)
                break
        if best:
            break
    if best is None:
        return []
    head_y, head_left, head_right = best
    header_words = sorted(word for line in layout if abs(line["y"] - head_y) <= 8 for word in _words_with_x(line))
    before = [right for left, right, _ in header_words if right <= head_left + 1]
    after = [left for left, right, _ in header_words if left >= head_right - 1]
    left = max(before) + 4 if before else head_left
    right = min(after) - 4 if after else head_right
    rows = {}
    row_texts = {}
    for line in layout:
        if line["y"] <= head_y + 8:
            continue
        row_texts.setdefault(round(line["y"] / 8), []).append(line["text"])
        for word_left, word_right, word in _words_with_x(line):
            middle = (word_left + word_right) / 2
            if left <= middle <= right:
                rows.setdefault(round(line["y"] / 8), []).append((word_left, word))
    if row_with:
        wanted_row = row_with.strip().lower()
        rows = {key: words for key, words in rows.items()
                if any(wanted_row in " ".join(row_texts.get(near, [])).lower() for near in (key - 1, key, key + 1))}
    return [" ".join(word for _, word in sorted(rows[key])) for key in sorted(rows)]


def clean_number(value: str) -> str:
    """An OCR'd number with look-alike letters ("483S236") -> digits, when
    that's what it is; otherwise unchanged."""
    fixed = value.translate(_DIGIT_LOOKALIKES).strip()
    return fixed if re.fullmatch(r"[0-9][0-9 ./-]*", fixed) and sum(ch.isdigit() for ch in value) >= len(value.strip()) / 2 else value.strip()


def _column_of(layout: list, anchor: dict) -> list[dict]:
    """The texts straight above / below `anchor` (overlapping it
    horizontally): the items of a popup list like the tours after F1."""
    left, right = anchor["x"], anchor["x"] + max(1, anchor.get("w") or len(anchor["text"]) * 14)
    column = []
    for line in layout:
        line_left, line_right = line["x"], line["x"] + max(1, line.get("w") or len(line["text"]) * 14)
        overlap = min(right, line_right) - max(left, line_left)
        if overlap >= 0.4 * min(right - left, line_right - line_left):
            column.append(line)
    items = []
    for line in sorted(column, key=lambda line: line["y"]):
        if items and abs(line["y"] - items[-1]["y"]) <= 8:
            items[-1]["text"] += " " + line["text"]
            items[-1]["highlighted"] = items[-1]["highlighted"] or bool(line.get("cursor"))
        else:
            items.append({"y": line["y"], "text": line["text"], "highlighted": bool(line.get("cursor"))})
    return items


def list_items(layout: list, target: str) -> list[dict]:
    """The list to move through, as rows {text, y, highlighted}. A popup
    list lies on top of a table (its own teal cell vs the table's teal row):
    then it's the column holding the value with a teal cell in it, or the
    column of a lone teal cell. Otherwise the screen's rows."""
    # An open popup list: its lone teal cursor cell. Then only that list
    # counts -- the unbroken run of rows around the cursor -- never the same
    # words elsewhere (a field above it showing the last choice,
    # "[OneDrive CSV", made the run steer at that field).
    cursor_rows = {}
    for line in layout:
        if line.get("cursor"):
            cursor_rows.setdefault(round(line["y"] / 8), []).append(line)
    lone = [lines[0] for lines in cursor_rows.values() if len(lines) == 1 and (lines[0].get("w") or 999) < 200]
    if lone and len(cursor_rows) > 1:
        # Several single teal cells (small markers like "(4)" or a page
        # counter too): a real popup is the longest list around one of them,
        # 5+ rows with that one teal row only.
        lists = [_contiguous_around_cursor(_column_of(layout, cell), cell["y"]) for cell in lone]
        popups = [items for items in lists if len(items) >= 5 and sum(1 for item in items if item["highlighted"]) == 1]
        if popups:
            return max(popups, key=len)
    for line in layout:
        if value_matches(line["text"], target):
            column = _column_of(layout, line)
            if any(item["highlighted"] for item in column):
                return column
    if lone and len(cursor_rows) > 1:
        return _column_of(layout, lone[0])
    return screen_rows(layout)


def _contiguous_around_cursor(items: list[dict], anchor_y: float | None = None) -> list[dict]:
    """The rows of one list: from the highlighted row up and down while the
    rows follow each other (a gap of more than about one and a half rows =
    something else on the screen above or below the list)."""
    # From the given cursor cell (a table row behind a popup is teal too).
    cursor = next((index for index, item in enumerate(items) if item["highlighted"] and (anchor_y is None or abs(item["y"] - anchor_y) <= 8)), None)
    if cursor is None or len(items) < 3:
        return items
    gaps = sorted(items[index + 1]["y"] - items[index]["y"] for index in range(len(items) - 1))
    step = gaps[len(gaps) // 2] or 26
    start, end = cursor, cursor
    while start > 0 and items[start]["y"] - items[start - 1]["y"] <= 1.6 * step:
        start -= 1
    while end < len(items) - 1 and items[end + 1]["y"] - items[end]["y"] <= 1.6 * step:
        end += 1
    return items[start:end + 1]


def row_key(row_text: str) -> str:
    """What to look for in a list row next time: its first number (the tour,
    e.g. "508"), else its first word."""
    number = re.search(r"(?<![\w.,])\d+(?![\w.,])", row_text)
    if number:
        return number.group()
    words = row_text.split()
    return words[0] if words else row_text


# ------------------------------------------------------------------ window
def _window_process(hwnd: int) -> int:
    process = wintypes.DWORD()
    ctypes.windll.user32.GetWindowThreadProcessId(wintypes.HWND(hwnd), ctypes.byref(process))
    return process.value


def find_uniware_window(title_part: str = DEFAULT_WINDOW_TITLE) -> int | None:
    found = []

    @ctypes.WINFUNCTYPE(wintypes.BOOL, wintypes.HWND, wintypes.LPARAM)
    def callback(hwnd, _):
        if user32.IsWindowVisible(hwnd) and title_part.lower() in screen_reader.window_title(hwnd).lower():
            found.append(hwnd)
        return True

    user32.EnumWindows(callback, 0)
    return found[0] if found else None


def bring_to_front(hwnd: int) -> bool:
    if user32.IsIconic(hwnd):
        user32.ShowWindow(hwnd, 9)  # SW_RESTORE
    # Windows only lets a process take focus right after input; a harmless
    # Alt tap counts as that.
    controller = keyboard.Controller()
    controller.press(keyboard.Key.alt)
    controller.release(keyboard.Key.alt)
    user32.SetForegroundWindow(hwnd)
    time.sleep(0.4)
    return screen_reader.foreground_window() == hwnd


# ------------------------------------------------------------------ other windows
# Uniware's Ctrl+P (print) doesn't save a file where we want it: it opens the
# data in a new window of another program. The run waits for that window,
# finds the file it shows and copies it as the dump.
def window_pid(hwnd: int) -> int:
    pid = wintypes.DWORD()
    user32.GetWindowThreadProcessId(hwnd, ctypes.byref(pid))
    return pid.value


def process_exe(pid: int) -> str:
    handle = ctypes.windll.kernel32.OpenProcess(0x1000, False, pid)  # PROCESS_QUERY_LIMITED_INFORMATION
    if not handle:
        return ""
    try:
        buffer = ctypes.create_unicode_buffer(1024)
        size = wintypes.DWORD(len(buffer))
        if ctypes.windll.kernel32.QueryFullProcessImageNameW(handle, 0, buffer, ctypes.byref(size)):
            return buffer.value
        return ""
    finally:
        ctypes.windll.kernel32.CloseHandle(handle)


def visible_windows() -> dict:
    """hwnd -> (pid, title) of every visible top-level window with a title."""
    found = {}

    @ctypes.WINFUNCTYPE(wintypes.BOOL, wintypes.HWND, wintypes.LPARAM)
    def callback(hwnd, _):
        if user32.IsWindowVisible(hwnd):
            title = screen_reader.window_title(hwnd)
            if title.strip():
                found[hwnd] = (window_pid(hwnd), title)
        return True

    user32.EnumWindows(callback, 0)
    return found


def process_command_line(pid: int) -> str:
    try:
        result = subprocess.run(
            ["powershell", "-NoProfile", "-Command", f"(Get-CimInstance Win32_Process -Filter 'ProcessId={int(pid)}').CommandLine"],
            capture_output=True, text=True, timeout=15, creationflags=0x08000000,  # CREATE_NO_WINDOW
        )
        return result.stdout.strip()
    except (OSError, subprocess.SubprocessError):
        return ""


DUMP_SEARCH_DEPTH = 4


def _search_dirs(extra: list[str]) -> list[Path]:
    home = Path.home()
    dirs = [Path(path) for path in extra if path]
    for name in ("TEMP", "TMP", "OneDrive", "OneDriveCommercial", "OneDriveConsumer"):
        if os.environ.get(name):
            dirs.append(Path(os.environ[name]))
    dirs += [home / "Downloads", home / "Documents", home / "Desktop", home / "AppData" / "Local" / "Temp"]
    unique = []
    for folder in dirs:
        if folder.is_dir() and folder not in unique:
            unique.append(folder)
    return unique


def _find_named_file(name: str, since: float, folders: list[Path]) -> Path | None:
    wanted = name.lower()
    best = None
    for folder in folders:
        base_depth = len(folder.parts)
        for root, dirnames, filenames in os.walk(folder):
            if len(Path(root).parts) - base_depth >= DUMP_SEARCH_DEPTH:
                dirnames[:] = []
            for filename in filenames:
                if filename.lower() != wanted:
                    continue
                path = Path(root) / filename
                try:
                    modified = path.stat().st_mtime
                except OSError:
                    continue
                if modified >= since - 300 and (best is None or modified > best[0]):
                    best = (modified, path)
    return best[1] if best else None


def window_file_name(title: str) -> str:
    """The file name in a window title ("table_export_data_1.xlsx  -
    Beveiligde weergave - Excel" -> "table_export_data_1.xlsx"); "" when the
    title has none ("Openen -Excel" while Excel is still starting)."""
    name = re.sub(r"\s*\[[^\]]*\]\s*$", "", title.split(" - ")[0]).strip(" *")
    return name if re.search(r"\.[A-Za-z0-9]{2,5}$", name) else ""


def dump_file_of_window(pid: int, title: str, since: float, extra_dirs: list[str], used: set | None = None) -> Path | None:
    """The file a just-opened window shows. The file name in its title wins
    ("dump.filipe....csv - Excel"): an Excel that was already running still
    has an older file on its command line. Without a name in the title: a
    file on the command line (e.g. notepad.exe "C:\\...\\dump.csv") written
    since the run started. Never a file this run already took (used)."""
    used = used or set()
    candidates, folders = [], []
    for quoted, bare in re.findall(r'"([^"]+)"|(\S+)', process_command_line(pid)):
        path = Path(quoted or bare)
        if path.suffix.lower() != ".exe" and path.is_file():
            if path.parent not in folders:
                folders.append(path.parent)  # e.g. \\10.1.21.4\Excel (Uniware's share)
            if str(path).lower() not in used:
                candidates.append(path)
    name = window_file_name(title)
    if name:
        for path in candidates:
            if path.name.lower() == name.lower():
                return path
        # The same folder as the file Excel was started with.
        for folder in folders:
            path = folder / name
            if path.is_file() and str(path).lower() not in used:
                return path
        if Path(name).is_file() and str(Path(name)).lower() not in used:
            return Path(name)
        found = _find_named_file(name, since, _search_dirs(extra_dirs))
        return found if found is not None and str(found).lower() not in used else None
    recent = [path for path in candidates if path.stat().st_mtime >= since - 2]
    return max(recent, key=lambda path: path.stat().st_mtime) if recent else None


# ------------------------------------------------------------------ keys
def parse_combo(combo: str):
    """"Shift+F5" -> ([Key.shift], Key.f5); "a" -> ([], "a")."""
    if len(combo) == 1:
        return [], combo
    parts = combo.split("+")
    if combo.endswith("+"):  # "Ctrl++": the key itself is "+"
        parts = parts[:-2] + ["+"]
    *modifier_names, base = parts
    modifiers = []
    for name in modifier_names:
        modifiers.append({"ctrl": keyboard.Key.ctrl, "alt": keyboard.Key.alt, "altgr": keyboard.Key.alt_gr, "shift": keyboard.Key.shift}[name.lower()])
    if len(base) == 1:
        if modifiers and base.isascii() and base.isalnum():
            # With Ctrl/Alt the physical key, not the character: pynput
            # types "P" as a Unicode character, which Ctrl doesn't combine
            # with -- Uniware never saw Ctrl+P / Ctrl+A.
            return modifiers, keyboard.KeyCode.from_vk(ord(base.upper()))
        return modifiers, base
    key_name = base.lower().replace(" ", "_")
    if key_name == "space":
        return modifiers, " "
    return modifiers, getattr(keyboard.Key, key_name)


def send_combo(controller: keyboard.Controller, combo: str):
    modifiers, base = parse_combo(combo)
    for modifier in modifiers:
        controller.press(modifier)
    controller.press(base)
    controller.release(base)
    for modifier in reversed(modifiers):
        controller.release(modifier)


# ------------------------------------------------------------------ recording
class ActionRecorder:
    """Records one named action from the user's own key presses in Uniware.
    Ctrl+Alt+C marks the current screen to be captured, Ctrl+Alt+Q stops."""

    def __init__(self, name: str, description: str = "", title_part: str = DEFAULT_WINDOW_TITLE):
        self.action = {
            "name": name.strip() or "New action",
            "slug": slugify(name),
            "description": description.strip(),
            "created": datetime.datetime.now().isoformat(timespec="seconds"),
            "start_screen": "",
            "steps": [],
        }
        self.title_part = title_part
        self.hwnd = None
        self.held = set()
        self.events: queue.Queue = queue.Queue()
        self.active = False
        self.status = ""
        self.selected = []
        self._listener = None
        self._thread = None
        self._open = None  # (step, pressed at) until its screen is known
        self.uniware_pid = 0
        self.known_windows: set = set()
        self.screen_id = ""
        self.layout: list = []
        self._typed: list[str] = []
        self._typed_at = 0.0
        self._last_printable_at = 0.0
        self._helper = recorder.Recorder.__new__(recorder.Recorder)  # for combo_name()
        self._helper.held = self.held

    def start(self):
        self.hwnd = find_uniware_window(self.title_part)
        if not self.hwnd:
            raise RuntimeError("Uniware window not found -- start Uniware (logged in) first.")
        bring_to_front(self.hwnd)
        image = screen_reader.capture_when_settled(self.hwnd)
        screen_id, title, selected, _, layout = recorder.describe_screen(image)
        self.screen_id, self.layout = screen_id, layout
        self.action["start_screen"] = screen_id
        self.action["start_title"] = title
        self.uniware_pid = window_pid(self.hwnd)
        self.known_windows = set(visible_windows())
        self.selected = selected
        self.active = True
        self.status = f"Recording -- start screen: {title}"
        self._listener = keyboard.Listener(on_press=self._on_press, on_release=self._on_release)
        self._listener.start()
        self._thread = threading.Thread(target=self._worker, daemon=True)
        self._thread.start()

    def stop(self) -> dict:
        self.events.put(("stop", None, time.time(), None))
        if self._thread:
            self._thread.join(timeout=10)
        if self._listener:
            self._listener.stop()
        self.active = False
        save_action(self.action)
        self.status = f"Saved: {self.action['name']} ({len(self.action['steps'])} steps)"
        return self.action

    def mark_capture(self):
        self.events.put(("capture", None, time.time(), None))

    def _on_press(self, key):
        if key in recorder.MODIFIERS:
            self.held.add(recorder.MODIFIERS[key])
            return
        vk = getattr(key, "vk", None)
        if {"Ctrl", "Alt"} <= self.held and vk in (0x43, 0x51):  # C, Q
            self.events.put(("capture" if vk == 0x43 else "stop", None, time.time(), None))
            return
        if not self.active or screen_reader.foreground_window() != self.hwnd:
            return
        name, printable = self._helper.combo_name(key)
        now = time.time()
        # The screen as it is when this key goes down is where the previous
        # step ended up: you only press the next key once Uniware has finished
        # loading, however long that took. (Inside a run of typed characters
        # no new picture is needed.)
        starts_something = not printable or now - self._last_printable_at > TYPING_GAP_SECONDS
        image = screen_reader.capture_window(self.hwnd) if starts_something else None
        self._last_printable_at = now if printable else 0.0
        self.events.put(("key", (name, printable), now, image))

    def _on_release(self, key):
        if key in recorder.MODIFIERS:
            self.held.discard(recorder.MODIFIERS[key])

    def _finish_open_step(self, image, at: float):
        """The open step (last key, or a run of typed text) ends on the screen
        in `image`; remember where it went and how long Uniware took."""
        screen_id, title, selected, _, layout = recorder.describe_screen(image)
        self.selected = selected
        self.screen_id = screen_id
        self.layout = layout
        if self._open is None:
            return
        step, started = self._open
        self._open = None
        waited = max(0.0, at - started)
        menu_before = step.pop("_menu_before", None)
        if menu_before and not is_menu(screen_id, layout):
            # An arrow that left the menu (Right into a program): the arrows
            # before it become one route by name, this key opens it.
            while self.action["steps"] and self.action["steps"][-1].get("key") in MENU_KEYS and self.action["steps"][-1].get("to_menu"):
                self.action["steps"].pop()
            step["menu"] = menu_before
        step.update({
            "to": screen_id, "to_title": title, "capture": False, "capture_pages": False,
            "recorded_wait": round(waited, 1),
            "to_menu": is_menu(screen_id, layout),
            "max_wait": default_max_wait(waited),
            "wait_text": "",
        })
        self.action["steps"].append(step)
        self.status = f"Step {len(self.action['steps'])}: {step.get('key') or 'type ' + repr(step['text'])} -> {title} (took {waited:.0f}s)"
        windows = visible_windows()
        opened = [(pid, window_title) for hwnd, (pid, window_title) in windows.items()
                  if hwnd not in self.known_windows and pid != self.uniware_pid]
        self.known_windows = set(windows)
        if opened and "key" in step:
            pid, window_title = opened[0]
            step["opens_window"] = {"title": window_title, "exe": Path(process_exe(pid)).name}
            self.status = f"Step {len(self.action['steps'])}: {step['key']} opened '{window_title}' -- the run takes that file as the dump"

    def _close_typing(self, at: float):
        if self._typed:
            typed = "".join(" " if name == "Space" else name for name in self._typed)
            self._open = ({"text": typed, "recorded_text": typed}, self._typed_at)
            self._typed = []

    def _worker(self):
        while True:
            kind, value, at, image = self.events.get()
            if kind in ("stop", "capture"):
                # Nothing is pressed now, so the screen you're looking at is
                # the finished one.
                self._close_typing(at)
                if self._open is not None:
                    self._finish_open_step(screen_reader.capture_when_settled(self.hwnd), at)
                if kind == "stop":
                    return
                if self.action["steps"]:
                    self.action["steps"][-1]["capture"] = True
                    self.status = f"Step {len(self.action['steps'])} will capture its screen"
                else:
                    self.action["capture_start"] = True
                    self.status = "The start screen will be captured"
                continue
            name, printable = value
            if image is not None:
                self._close_typing(at)
                self._finish_open_step(image, at)
            if printable:
                if not self._typed:
                    self._typed_at = at
                self._typed.append(name)
                continue
            step = {"key": name}
            if name in MENU_KEYS and is_menu(self.screen_id, self.layout):
                # Arrow on a menu: if it turns out to leave the menu (Right
                # on "Inklaring bedrijf" opens the program), it becomes the
                # route step -- see _finish_open_step.
                step["_menu_before"] = menu_path(self.layout)
            if name not in MENU_KEYS and is_menu(self.screen_id, self.layout):
                # A key that leaves (or acts in) a menu: the arrows that led
                # here become one route by name, e.g. Expeditie > Inklaring
                # bedrijf -- replay steers there from wherever the menu is.
                steps = self.action["steps"]
                while steps and steps[-1].get("key") in MENU_KEYS and steps[-1].get("to_menu"):
                    steps.pop()
                step["menu"] = menu_path(self.layout)
                self.status = f"Menu route: {' > '.join(step['menu'])}, then {name}"
            elif name == "Enter":
                # Arrowed through a list (e.g. the tours after F1) and then
                # Enter: the row's position changes from day to day, so
                # replay looks the row up by its value instead of pressing
                # Down the same number of times.
                steps = self.action["steps"]
                arrows = 0
                while steps and steps[-1].get("key") in ARROW_KEYS and steps[-1].get("to") == self.screen_id:
                    steps.pop()
                    arrows += 1
                highlighted = [row for row in screen_rows(self.layout) if row["highlighted"]]
                if arrows and highlighted:
                    step["pick"] = row_key(highlighted[0]["text"])
                    step["pick_row"] = highlighted[0]["text"]
                    self.status = f"List: will look for the row with '{step['pick']}' ({highlighted[0]['text'][:50]}), not count {arrows} key(s) -- change it under Edit if needed"
                # Plain Enter in a list (no arrows): on whatever row is
                # selected then -- its text (an order, a customer) differs
                # every day, so it's never looked for.
            self._open = (step, at)


# ------------------------------------------------------------------ replay
class ActionRun:
    """Replays an action. After every step it waits until Uniware really shows
    the recorded screen (and the step's "wait until" text) and has stopped
    changing -- up to the step's max wait -- and stops on anything else."""

    def __init__(self, action: dict, values: dict, title_part: str = DEFAULT_WINDOW_TITLE):
        self.action = action
        self.values = {key: str(value) for key, value in (values or {}).items()}
        self.title_part = title_part
        self.log: list[str] = []
        self.active = False
        self.stop_requested = False
        self.result = ""
        self.output_files: list[str] = []
        self.exported_files: list[str] = []
        self.dump_parts: list[Path] = []  # "join the dumps": one per Ctrl+P
        self.used_dump_files: set[str] = set()  # never the same file twice
        # The collected columns waiting for the dump that only arrives at the
        # end (Uniware's own folder dump, "dump.filipe....csv").
        self.pending_dump_merge = None
        self.collected_csv_paths: set[str] = set()
        self.captures: list[dict] = []
        self.started = datetime.datetime.now()

    def say(self, message: str):
        line = f"{datetime.datetime.now():%H:%M:%S} {message}"
        self.log.append(line)
        # The whole log in a file too (the dashboard shows only the last lines).
        try:
            OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)
            with (OUTPUTS_DIR / f"{self.action.get('slug', 'action')}_{self.started:%Y%m%d_%H%M%S}.log").open("a", encoding="utf-8") as handle:
                handle.write(line + "\n")
        except (OSError, AttributeError):
            pass

    def _check(self, hwnd):
        if self.stop_requested:
            raise RuntimeError("Stopped by you.")
        front = screen_reader.foreground_window()
        if front != hwnd:
            if _window_process(front) == _window_process(hwnd):
                title = screen_reader.window_title(front)
                raise RuntimeError(f"Uniware opened another window ('{title}' -- a message or question?) -- stopped for safety.")
            raise RuntimeError("Uniware lost focus (another window was clicked) -- stopped for safety.")

    def _read(self, hwnd, timeout: float = 3.0):
        image = screen_reader.capture_when_settled(hwnd, timeout=timeout)
        return (image, *recorder.describe_screen(image))

    def _screenshot(self, hwnd):
        return screen_reader.capture_window(hwnd)

    def _describe(self, image):
        return recorder.describe_screen(image)

    def _read_settled(self, hwnd, timeout: float = 3.0):
        """Like _read_stable, but quicker: screenshots only (no text reading)
        until two in a row 0.25 s apart are the same, then read once."""
        deadline = time.time() + timeout
        image = self._screenshot(hwnd)
        while time.time() < deadline:
            time.sleep(0.25)
            shot = self._screenshot(hwnd)
            if not screens_differ(image, shot):
                image = shot
                break
            image = shot
        return (image, *self._describe(image))

    def _wait_for(self, hwnd, step: dict, label: str):
        """Read the screen until it is the step's screen, contains its "wait
        until" text and two reads in a row are the same (finished loading),
        for at most the step's max wait. Returns the reading."""
        expected = step["to"]
        expected_title = step.get("to_title") or expected
        wait_text = (step.get("wait_text") or "").strip()
        max_wait = float(step.get("max_wait") or default_max_wait(float(step.get("recorded_wait") or 0)))
        started = time.time()
        next_note = started + 5
        previous_text = None
        while True:
            self._check(hwnd)
            image, screen_id, title, selected, text, layout = self._read(hwnd)
            # The top of the screen too (a popup such as "Search" sits there).
            joined = "\n".join([title, *text]).lower()
            # No screen name (a popup without a stable one): only the
            # "wait until" text decides.
            if (not expected or same_screen(expected, expected_title, screen_id, title)) and wait_text.lower() in joined and not BUSY_RE.search(joined):
                if text == previous_text:
                    return screen_id, title, selected, text, layout
                previous_text = text
                continue  # read once more: is it still drawing?
            previous_text = None
            waited = time.time() - started
            if waited > max_wait:
                self._save_problem_screen(image)
                looking = f"'{expected_title}'" + (f" showing '{wait_text}'" if wait_text else "")
                raise RuntimeError(f"{label}: waited {max_wait:.0f}s for {looking}, but Uniware shows '{title}'. Stopped (screenshot in outputs).")
            if time.time() > next_note:
                self.say(f"{label}: Uniware still busy, waiting for {expected_title}" + (f" / '{wait_text}'" if wait_text else "") + f" ({waited:.0f}s of max {max_wait:.0f}s)")
                next_note = time.time() + 10
            time.sleep(0.5)

    def _read_column_value(self, hwnd, step: dict, label: str, layout: list):
        """Read a column (e.g. invoiceRes) off the screen into a field for
        the dump file name. The value can appear a moment after the key
        (Shift+F5 makes the invoice), so it waits -- up to the step's max
        wait -- until the column is filled and two separate screenshots
        read the same: a misread number must never end up in a file name
        that SnappySjaak takes as the invoice number."""
        column, field = step["read_column"], step["read_field"]
        max_wait = float(step.get("max_wait") or 60)
        started = time.time()
        next_note = started + 5
        previous = None
        while True:
            self._check(hwnd)
            values = [clean_number(value) for value in read_column(layout, column, step.get("read_row_with") or "")]
            if values and values == previous:
                break
            previous = values or None
            waited = time.time() - started
            if waited > max_wait:
                if not values:
                    raise RuntimeError(f"{label}: column '{column}' stayed empty for {max_wait:.0f}s (or isn't on this screen) -- nothing to put in {{{field}}}.")
                raise RuntimeError(f"{label}: column '{column}' kept reading differently ({', '.join(dict.fromkeys(values))}). Stopped to be safe.")
            if time.time() > next_note:
                self.say(f"{label}: waiting for '{column}' to be filled ({waited:.0f}s of max {max_wait:.0f}s)")
                next_note = time.time() + 10
            time.sleep(0.5)
            _, _, _, _, _, layout = self._read(hwnd, timeout=1.5)
        distinct = list(dict.fromkeys(values))
        # SnappySjaak takes the LAST number in the dump's file name as the
        # invoice number (ukdocsInvoiceNumberFromFileName): it has to be one
        # value, digits only -- otherwise it would quietly pick the wrong one.
        if len(distinct) > 1:
            raise RuntimeError(f"{label}: '{column}' has {len(distinct)} different values ({', '.join(distinct)}) -- SnappySjaak expects one invoice number per file. Stopped.")
        if field == "invoiceres" and not distinct[0].isdigit():
            raise RuntimeError(f"{label}: '{column}' reads '{distinct[0]}', not just digits -- SnappySjaak would take only part of it as the invoice number. Stopped.")
        self.values[field] = distinct[0]
        self.say(f"{label}: {column} = {self.values[field]} -> {{{field}}}")

    def _take_window_dump(self, hwnd, windows_before: dict, step: dict, label: str):
        """The key opened the data in another program's window (Uniware's
        Ctrl+P): wait for that window, find its file, copy it as the dump
        (renamed, to the dump folder) and close the window."""
        uniware_pid = window_pid(hwnd)
        max_wait = float(step.get("max_wait") or 60)
        started = time.time()
        window = None
        while window is None:
            if self.stop_requested:
                raise RuntimeError("Stopped by you.")
            for other, (pid, title) in visible_windows().items():
                if other not in windows_before and pid != uniware_pid:
                    window = (other, pid, title)
                    break
            if window is None:
                if time.time() - started > max_wait:
                    raise RuntimeError(f"{label}: no new window opened within {max_wait:.0f}s after {step['key']}.")
                time.sleep(0.5)
        other, pid, title = window
        time.sleep(1.5)  # let the program put the file name in its title

        def program_windows():
            # The program's new windows: Excel first shows "Openen", then
            # the workbook in another window.
            return {handle: info[1] for handle, info in visible_windows().items() if handle not in windows_before and info[0] == pid}

        def shown_title():
            titles = [screen_reader.window_title(handle) or text for handle, text in program_windows().items()]
            return next((text for text in titles if window_file_name(text)), titles[0] if titles else title)

        title = shown_title()
        self.say(f"{label}: {step['key']} opened '{title}' ({Path(process_exe(pid)).name})")
        source = None
        extra = [self.action.get("uniware_dump_dir") or "", self.action.get("export_dir") or ""]
        while source is None:
            title = shown_title()
            source = dump_file_of_window(pid, title, self.started.timestamp(), extra, self.used_dump_files)
            if source is None:
                if time.time() - started > max_wait + 30:
                    raise RuntimeError(f"{label}: could not find the file shown in '{title}'. Tell which folder it is in under Edit (Uniware saves its dump in).")
                time.sleep(1)
        previous_size = -1
        while True:  # still being written?
            size = source.stat().st_size
            if size > 0 and size == previous_size:
                break
            previous_size = size
            time.sleep(1)
        self.used_dump_files.add(str(source).lower())
        if self.action.get("join_dumps"):
            # Several dumps in one run (e.g. tour 508: one per company) become
            # one file at the end -- named then, so a value read later (the
            # LTD invoiceRes) can be in the name.
            DUMP_PARTS_DIR.mkdir(parents=True, exist_ok=True)
            part = DUMP_PARTS_DIR / f"{self.action['slug']}_{self.started:%Y%m%d_%H%M%S}_part{len(self.dump_parts) + 1}{source.suffix or '.txt'}"
            part.write_bytes(source.read_bytes())
            self.dump_parts.append(part)
            self.say(f"{label}: dump {len(self.dump_parts)} kept ({source.name})")
            window_steps = [item for item in self.action.get("steps", []) if item.get("opens_window")]
            if window_steps and step is window_steps[-1]:
                # The last dump of the action: join now, not only at the
                # end -- a problem on the way back must not lose the file.
                self._join_dump_parts()
                self.dump_parts = []
        else:
            target_dir = self._export_dir()
            target = target_dir / f"{self.export_name()}{source.suffix or '.txt'}"
            temp = target_dir / f"~{target.name}.tmp"
            temp.write_bytes(source.read_bytes())  # copy: the program may still hold it open
            temp.replace(target)
            self.exported_files.append(str(target))
            self.say(f"{label}: dump {source} -> {target}")
        if self.action.get("close_dump_window", True):
            # Close every new window of that program -- also one that only
            # appears now (the workbook after Excel's "Openen") -- until none
            # is left for 2 seconds.
            quiet_since = None
            close_started = time.time()
            while time.time() - close_started < 20:
                windows = program_windows()
                for handle in windows:
                    user32.PostMessageW(handle, 0x0010, 0, 0)  # WM_CLOSE
                if windows:
                    quiet_since = None
                elif quiet_since is None:
                    quiet_since = time.time()
                elif time.time() - quiet_since >= 2:
                    break
                time.sleep(0.5)

    def _export_dir(self) -> Path:
        target_dir = Path((self.action.get("export_dir") or "").strip() or (self.action.get("uniware_dump_dir") or "").strip() or OUTPUTS_DIR)
        target_dir.mkdir(parents=True, exist_ok=True)
        return target_dir

    def _join_dump_parts(self):
        """The run's dumps as one file: the first one's header, then the rows
        of all of them (same columns required; itemNumber numbered on)."""
        suffixes = {part.suffix.lower() for part in self.dump_parts}
        if len(suffixes) != 1:
            raise RuntimeError(f"The dumps are different kinds of file ({', '.join(sorted(suffixes))}) -- can't join them.")
        suffix = suffixes.pop()
        target = self._export_dir() / f"{self.export_name()}{suffix}"
        temp = target.parent / f"~{target.name}.tmp"
        counts = []
        if suffix == ".xlsx":
            import openpyxl  # only needed here
            workbook = openpyxl.load_workbook(self.dump_parts[0])
            sheet = workbook.active
            header = [cell.value for cell in sheet[1]]
            counts.append(sheet.max_row - 1)
            number_column = header.index("itemNumber") if "itemNumber" in header else None
            last_number = max((row[number_column] for row in sheet.iter_rows(min_row=2, values_only=True)
                               if number_column is not None and isinstance(row[number_column], (int, float))), default=0)
            for part in self.dump_parts[1:]:
                other_book = openpyxl.load_workbook(part, read_only=True)
                rows = list(other_book.active.iter_rows(values_only=True))
                other_book.close()  # read-only keeps the file open otherwise
                if not rows or _trimmed(rows[0]) != _trimmed(header):
                    raise RuntimeError(f"Dump {part.name} has other columns than the first one -- not joined (the parts are in {DUMP_PARTS_DIR}).")
                for row in rows[1:]:
                    if not any(value not in (None, "") for value in row):
                        continue
                    row = list(row)
                    if number_column is not None and isinstance(row[number_column], (int, float)):
                        last_number += 1
                        row[number_column] = last_number
                    sheet.append(row)
                counts.append(len(rows) - 1)
            workbook.save(temp)
            _relative_sheet_paths(temp)
        else:
            lines_out, header = [], None
            for part in self.dump_parts:
                lines = part.read_text(encoding="utf-8-sig", errors="replace").splitlines()
                if not lines:
                    continue
                if header is None:
                    header = lines[0]
                    lines_out.append(header)
                elif lines[0] != header:
                    raise RuntimeError(f"Dump {part.name} has other columns than the first one -- not joined (the parts are in {DUMP_PARTS_DIR}).")
                lines_out.extend(line for line in lines[1:] if line.strip())
                counts.append(len(lines) - 1)
            temp.write_text("\n".join(lines_out) + "\n", encoding="utf-8")
        temp.replace(target)
        self.exported_files.append(str(target))
        self.say(f"{len(self.dump_parts)} dumps joined ({' + '.join(str(count) for count in counts)} rows) -> {target}")

    def _back_to_menu(self, hwnd, controller, reading):
        """From inside a program, press the back key (Esc by default, Edit:
        "Back key") until a menu shows -- at most 8 times, and never once a
        menu is there (Esc on the main menu could close Uniware)."""
        back_key = self.action.get("back_key") or "Esc"
        screen_id, title, selected, text, layout = reading
        unchanged = 0
        for _ in range(8):
            if is_menu(screen_id, layout):
                return screen_id, title, selected, text, layout
            self._check(hwnd)
            self.say(f"Not in a menu ('{title}'): pressing {back_key}")
            send_combo(controller, back_key)
            _, new_id, title, selected, new_text, layout = self._read(hwnd)
            unchanged = unchanged + 1 if new_text == text else 0
            if unchanged >= 2:
                break
            screen_id, text = new_id, new_text
        raise RuntimeError(f"Could not get back to the menu from '{title}' with {back_key}. Go to the menu yourself and start again (or set the Back key under Edit).")

    def _go_menu(self, hwnd, controller, path: list, label: str):
        """Steer the menu to `path` (blue item per level) from wherever it is:
        Left back to the first level that's wrong, Up/Down within the level
        (counting rows from the blue item), Right into the next level --
        reading the screen after every move."""
        path = [item for item in path if item]
        if not path:
            return
        route = " > ".join(path)
        last_depth = None
        for _ in range(80):
            self._check(hwnd)
            image, screen_id, title, _, _, layout = self._read(hwnd, timeout=1.5)
            if not is_menu(screen_id, layout):
                self._save_problem_screen(image)
                raise RuntimeError(f"{label}: expected a menu on the way to {route}, but Uniware shows '{title}'. Stopped (screenshot in outputs).")
            levels = menu_levels(layout)
            active = max(index for index, level in enumerate(levels) if level["selected"] is not None)
            current = [levels[index]["items"][levels[index]["selected"]] if levels[index]["selected"] is not None else "" for index in range(active + 1)]
            wrong = next((index for index in range(min(active + 1, len(path))) if not same_item(current[index], path[index])), None)
            if wrong is None and active == len(path) - 1:
                self.say(f"{label}: menu at {route}")
                return
            if wrong is not None and wrong < active:
                self._press(controller, "Left", active - wrong)
            elif wrong is None and active > len(path) - 1:
                self._press(controller, "Left", active - (len(path) - 1))
            elif wrong is None:
                if last_depth is not None and last_depth == active:
                    self._save_problem_screen(image)
                    raise RuntimeError(f"{label}: Right on '{current[active]}' did not open its submenu (looking for {route}). Stopped (screenshot in outputs).")
                last_depth = active
                self._press(controller, "Right", 1)
                continue
            else:
                items = levels[active]["items"]
                target = next((index for index, item in enumerate(items) if same_item(item, path[active])), None)
                if target is None:
                    self._save_problem_screen(image)
                    raise RuntimeError(f"{label}: '{path[active]}' is not in this menu level ({', '.join(items[:8])}...). Stopped (screenshot in outputs).")
                distance = target - levels[active]["selected"]
                self._press(controller, "Down" if distance > 0 else "Up", abs(distance))
            last_depth = None
        raise RuntimeError(f"{label}: could not reach {route} in the menu. Stopped.")

    def _read_stable(self, hwnd):
        """Read until two screenshots in a row give the same text (the info
        panel next to a list updates a moment after the cursor moves)."""
        previous = None
        reading = None
        for _ in range(8):
            reading = self._read(hwnd, timeout=1.5)
            if previous is not None and reading[4] == previous:
                return reading
            previous = reading[4]
        return reading

    def _run_list_loop(self, hwnd, controller, loop: dict, label: str):
        """Go through a list row by row: when the screen shows one of the
        match words (e.g. the order type "Accessoire" next to the list), press
        the match key (Delete: out of the tour / Enter: into it), otherwise
        the next key (Down). Only acts when the type is clearly readable --
        one of the match or keep words, read the same twice; an unreadable
        row is skipped and logged. Stops at the end of the list (the cursor
        row no longer changes) or after max_rows. dry_run: only logs what it
        would do and moves on with the next key."""
        match_words = [word.strip() for word in loop.get("match", []) if word.strip()]
        keep_words = [word.strip() for word in loop.get("keep", []) if word.strip()]
        match_key = loop.get("match_key") or "Delete"
        next_key = loop.get("next_key") or "Down"
        dry_run = bool(loop.get("dry_run"))
        row_stays = bool(loop.get("row_stays"))
        only_rows_with = str(loop.get("only_rows_with") or "").strip()
        max_rows = max(1, min(2000, int(loop.get("max_rows") or 500)))
        collect_into = str(loop.get("collect_into") or "").strip()
        use_list = str(loop.get("use_list") or "").strip()
        use_label = str(loop.get("use_label") or "").strip()
        if not match_words and not collect_into and not use_list:
            raise RuntimeError(f"{label}: the list loop has no match words (Edit).")
        if use_list and not use_label:
            raise RuntimeError(f"{label}: the list loop acts on day list '{use_list}' but has no label (e.g. 516).")
        collected = {"name": collect_into, "items": []} if collect_into else None
        # Collect mode reads every list row on screen, not only the cursor row
        # (OCR misreads digits on the teal cursor: 4841449 -> "4 41449"); a
        # number counts once it was read the same way twice.
        sightings: dict[str, dict] = {}
        row_key = list_row_lot if loop.get("collect_key") == "lot" else list_row_order
        value_label = str(loop.get("collect_value") or "").strip()
        # The value (e.g. EPP) belongs to the cursor row only: key -> values read.
        row_values: dict[str, list[str]] = {}
        keys_by_position: dict[int, str] = {}  # screen row -> key, from the previous look
        popup_keys = str(loop.get("row_popup_keys") or "").strip()
        close_keys = str(loop.get("row_close_keys") or "").strip() or "Esc"
        row_popups: dict[str, list[str]] = {}  # key -> the popup's lines
        cursor_keys_done: set[str] = set()  # rows the cursor was on
        cursor_key = None
        last_others = None

        def note_rows(layout, joined):
            for screen_row in screen_rows(layout):
                order = row_key(screen_row["text"])
                if not order:
                    continue
                seen = sightings.setdefault(order, {"count": 0, "row": screen_row["text"][:80], "type": "", "brief": ""})
                seen["count"] += 1
                seen["brief"] = seen["brief"] or list_row_brief(screen_row["text"]) or ""
                if screen_row["highlighted"] and not seen["type"]:
                    seen["type"] = found(match_words, joined) or found(keep_words, joined) or ""
        list_data = load_day_list(use_list) if use_list else None
        # Acting on a day list: the orders still to do -- after the last one
        # the loop ends right away (a Down more can open another screen).
        todo = {item["key"] for item in list_data.get("items", []) if item.get("label") == use_label and not item.get("done")} if list_data is not None else set()
        if list_data is not None and not todo:
            self.say(f"{label}: no orders with result {use_label} left to do in '{use_list}'.")
            return
        if list_data is not None:
            self.say(f"{label}: {len(todo)} order(s) with result {use_label} to do: {', '.join(sorted(todo))}")

        def finished(entry) -> bool:
            if list_data is None or entry is None:
                return False
            todo.discard(entry["key"])
            if todo:
                return False
            self.say(f"{label}: all orders with result {use_label} done -- end of the loop.")
            return True

        def found(words, text):
            for word in words:
                if re.search(rf"(?<![\w]){re.escape(word)}(?![\w])", text, re.IGNORECASE):
                    return word
            return None

        def cursor_row(layout):
            rows = screen_rows(layout)
            highlighted = [row["text"] for row in rows if row["highlighted"]]
            return highlighted[0] if highlighted else ""

        def row_number_key(text):
            """The row's order number (6-8 digits, look-alike letters read
            as digits) -- neighbouring orders are up to ~87% alike as text."""
            return row_order_number(text)

        def same_row(left, right):
            if not left or not right:
                return False
            left_key, right_key = row_number_key(left), row_number_key(right)
            if left_key and right_key:
                return left_key == right_key
            return difflib.SequenceMatcher(None, left, right).ratio() >= 0.95

        counts = {"matched": 0, "kept": 0, "unreadable": 0}
        loop_started = time.time()
        last_row, last_action = None, None
        last_cursor_y = None
        handled = []
        unchanged = 0
        list_screen = None
        last_image = None
        for row_number in range(1, max_rows + 1):
            self._check(hwnd)
            if collected is not None and last_image is not None:
                # After the key: first wait until the screen changed (a slow
                # screen read too early looks like the end -- the next key
                # again would skip a row); at the end of the list it doesn't.
                deadline = time.time() + 2
                while time.time() < deadline and not screens_differ(last_image, self._screenshot(hwnd)):
                    time.sleep(0.1)
            image, screen_id, screen_title, _, text, layout = (self._read_settled(hwnd) if collected is not None else self._read_stable(hwnd))
            joined = "\n".join(text)
            row = cursor_row(layout)
            if collected is not None and row_number % 50 == 0:
                self.say(f"{label}: {row_number} rows, {(time.time() - loop_started) / row_number:.1f} s per row")
            # Never go on outside the list: a key that left its screen (into
            # a menu, another program) would otherwise be pressed there.
            if list_screen is None:
                list_screen = (screen_id, screen_title)
            elif not same_screen(list_screen[0], list_screen[1], screen_id, screen_title) or difflib.SequenceMatcher(
                    None, re.sub(r"[^a-z]+", "", list_screen[1].lower()), re.sub(r"[^a-z]+", "", (screen_title or "").lower())).ratio() < 0.85:
                # (Same program, other part: "Transport afspraken (In te
                # delen" -> "Transport afspraken" is not the list any more.)
                raise RuntimeError(f"{label}: Uniware left the list ('{list_screen[1]}' -> '{screen_title}') -- stopped to be safe.")
            if collected is not None:
                note_rows(layout, joined)
                rows_now = screen_rows(layout)
                cursor_entry = next((entry for entry in rows_now if entry["highlighted"]), None)
                cursor_key = row_key(row)
                if not cursor_key and cursor_entry is not None:
                    # The teal cursor row misread: the same screen row on
                    # the previous look (not under the cursor then) -- not
                    # one the cursor was on already (the list jumped) --
                    # else the row read again, enlarged.
                    cursor_key = keys_by_position.get(round(cursor_entry["y"] / 8))
                    if cursor_key in cursor_keys_done:
                        cursor_key = None
                    cursor_key = cursor_key or read_cursor_row_key(image, cursor_entry, row_key)
                if value_label:
                    value = read_screen_value(value_label, [line.get("text", "") for line in layout])
                    if cursor_key and value:
                        row_values.setdefault(cursor_key, []).append(value)
                keys_by_position.clear()
                for entry in rows_now:
                    key = row_key(entry["text"]) if not entry["highlighted"] else None
                    if key:
                        keys_by_position[round(entry["y"] / 8)] = key
            if last_action == "match" and same_row(row, last_row):
                raise RuntimeError(f"{label}: '{last_row[:50]}' is still there after {match_key} -- does Uniware ask for a confirmation? Stopped to be safe.")
            # Where the cursor row is on the screen: rows that look alike
            # (four "PROTEA CYNAROIDES ... Afgemeld" one after another) are
            # still new rows when the cursor moved down to them.
            cursor_entry_now = next((entry for entry in screen_rows(layout) if entry["highlighted"]), None)
            cursor_y = round(cursor_entry_now["y"] / 8) if cursor_entry_now else None
            cursor_moved = cursor_y is not None and last_cursor_y is not None and cursor_y != last_cursor_y
            # At the bottom of the screen the cursor stays and the list
            # scrolls under it (three "LI OT ZAMBESI" rows alike but their
            # lot): the other rows changed -- a new row too.
            others = tuple(row_key(entry["text"]) or row_order_number(entry["text"]) or "" for entry in screen_rows(layout) if not entry["highlighted"])
            cursor_moved = cursor_moved or (last_others is not None and others != last_others)
            last_cursor_y, last_others = cursor_y, others
            if last_action == "next" and same_row(row, last_row) and not cursor_moved:
                # No new order after the key: once can be a slow screen --
                # LIST_LOOP_END_TRIES times in a row is the end of the list.
                unchanged += 1
                if unchanged >= LIST_LOOP_END_TRIES:
                    self.say(f"{label}: no new row after {LIST_LOOP_END_TRIES} tries -- end of the list.")
                    break
                for part in list_loop_keys(next_key):
                    send_combo(controller, part)
                time.sleep(0.6)
                continue
            unchanged = 0
            if last_action == "match" and any(same_row(row, done) for done in handled):
                # The last row was taken out: the cursor went back to one
                # already handled -- move on (and find the end) quietly.
                for part in list_loop_keys(next_key):
                    send_combo(controller, part)
                last_row, last_action = row, "next"
                continue
            handled.append(row)
            if only_rows_with and not value_matches(row, only_rows_with):
                self.say(f"{label}: {row[:60] or f'row {row_number}'}: no '{only_rows_with}' -> passed by")
                for part in list_loop_keys(next_key):
                    send_combo(controller, part)
                last_row, last_action = row, "next"
                time.sleep(0.3)
                continue
            if collected is not None:
                # Collect mode: the rows are noted from the whole screen
                # (note_rows); only the next key is pressed -- after opening
                # and reading the row's popup, when one is set.
                if popup_keys:
                    # Every lot, also one with a brief number in Info: that
                    # can be the last digits of its EPP, not the real one.
                    self._read_row_popup(hwnd, controller, cursor_key, image, popup_keys, close_keys, row_popups, label)
                if cursor_key:
                    cursor_keys_done.add(cursor_key)
                for part in list_loop_keys(next_key):
                    send_combo(controller, part)
                last_row, last_action = row, "next"
                continue
            entry = None
            if list_data is not None:
                # Act on the day list: the row's order number decides.
                order = row_order_number(row)
                entry = day_list_item(list_data, order) if order else None
                if entry is None:
                    # The teal cursor row's digits are often misread ("4 41449"):
                    # the one day-list order that reads almost the same.
                    digits = re.sub(r"\D", "", " ".join(row.translate(_DIGIT_LOOKALIKES).split()[1:3]))
                    close = [item for item in list_data.get("items", [])
                             if len(digits) >= 5 and difflib.SequenceMatcher(None, digits, item["key"]).ratio() >= 0.85]
                    entry = close[0] if len(close) == 1 else None
                matched = f"{use_list}: {use_label}" if entry and entry.get("label") == use_label else None
                kept = None if matched else (f"{use_list}: {entry.get('label') or 'no label'}" if entry else "not on the day list")
            else:
                matched = found(match_words, joined)
                kept = found(keep_words, joined) if not matched else None
            if list_data is None and not matched and not kept:
                # Read once more before giving up on this row.
                _, _, _, _, text, layout = self._read_stable(hwnd)
                joined = "\n".join(text)
                matched = found(match_words, joined)
                kept = found(keep_words, joined) if not matched else None
            short = row[:60] or f"row {row_number}"
            if matched and not dry_run and row_stays:
                counts["matched"] += 1
                self.say(f"{label}: {short}: {matched} -> {match_key}")
                if list_data is not None and entry is not None:
                    entry["done"] = datetime.datetime.now().isoformat(timespec="seconds")
                    save_day_list(list_data)
                for index, part in enumerate(list_loop_keys(match_key)):
                    if index:
                        self._read_stable(hwnd)
                        self._check(hwnd)
                    send_combo(controller, part)
                # The row stays: it must have changed (e.g. its tour number
                # cleared), then on to the next row.
                _, _, _, _, _, after_layout = self._read_stable(hwnd)
                after = cursor_row(after_layout)
                if after == row:
                    raise RuntimeError(f"{label}: '{row[:50]}' did not change after {match_key}. Stopped to be safe.")
                if finished(entry):
                    break
                for part in list_loop_keys(next_key):
                    send_combo(controller, part)
                last_row, last_action = after, "next"
            elif matched and not dry_run:
                counts["matched"] += 1
                self.say(f"{label}: {short}: {matched} -> {match_key}")
                if list_data is not None and entry is not None:
                    entry["done"] = datetime.datetime.now().isoformat(timespec="seconds")
                    save_day_list(list_data)
                # The key(s) one after another, the screen settled in
                # between (a confirmation question has to be there before
                # it's answered).
                for index, part in enumerate(list_loop_keys(match_key)):
                    if index:
                        self._read_stable(hwnd)
                        self._check(hwnd)
                    send_combo(controller, part)
                if finished(entry):
                    break
                last_row, last_action = row, "match"
            else:
                if matched:
                    counts["matched"] += 1
                    self.say(f"{label}: {short}: {matched} -> would press {match_key} (trial run)")
                    if finished(entry):
                        break
                elif kept:
                    counts["kept"] += 1
                    self.say(f"{label}: {short}: {kept} -> keep")
                else:
                    counts["unreadable"] += 1
                    self.say(f"{label}: {short}: type not readable -> skipped (check this one)")
                for part in list_loop_keys(next_key):
                    send_combo(controller, part)
                last_row, last_action = row, "next"
            time.sleep(0.3)
        else:
            self.say(f"{label}: stopped after {max_rows} rows (the most this loop may do).")
        if collected is not None:
            # Cursor-row keys are read from the teal row (digits misread more
            # often there): one that isn't among the well-read keys goes to
            # the one well-read key it is nearly the same as.
            good_keys = [order for order, seen in sightings.items() if seen["count"] >= 2]
            values_by_key: dict[str, str] = {}
            for cursor_key, values in row_values.items():
                key = cursor_key if cursor_key in good_keys else None
                if key is None:
                    close = [order for order in good_keys if difflib.SequenceMatcher(None, cursor_key, order).ratio() >= 0.85]
                    key = close[0] if len(close) == 1 else None
                if key:
                    # The value read most often for that row (it is read on every look).
                    values_by_key[key] = max(set(values), key=values.count)
            for order, seen in sightings.items():
                if seen["count"] < 2:
                    self.say(f"{label}: {order} read only once ('{seen['row'][:50]}') -- not collected, check by hand")
                    continue
                item = {"key": order, "row": seen["row"], "type": seen["type"], "label": "", "note": "", "warning": False, "brief": seen.get("brief", "")}
                if value_label:
                    # Not every lot has one (e.g. no EPP): left empty.
                    item["value"] = values_by_key.get(order, "")
                    if not item["value"]:
                        item["note"] = f"no {value_label}"
                if popup_keys:
                    item["info"] = " | ".join(row_popups.get(order, []))
                    if not item["info"]:
                        item["note"] = (item.get("note", "") + "; " if item.get("note") else "") + "popup not read"
                collected["items"].append(item)
                self.say(f"{label}: collected {order} (Info {item['brief'] or 'empty'})" + (f" {value_label} {item.get('value') or 'none'}" if value_label else f" ({seen['type'] or 'type ?'})")
                         + (f", popup {len(row_popups.get(order, []))} line(s)" if popup_keys else ""))
            save_day_list(collected)
            self.say(f"{label}: {len(collected['items'])} row(s) collected into day list '{collect_into}'.")
            if value_label or popup_keys:
                self._write_collected_csv(collected, value_label or "value", label)
            return
        self.say(f"{label}: list done -- {counts['matched']} {'would be ' if dry_run else ''}{match_key}, {counts['kept']} kept, {counts['unreadable']} not readable.")

    def _read_row_popup(self, hwnd, controller, key, image, popup_keys, close_keys, row_popups, label):
        """Open the cursor row's popup (e.g. ? -- the supplier info), keep the
        lines inside it (the part of the screen that changed), close it and
        wait until the list is back."""
        if not key or key in row_popups:
            return
        for part in list_loop_keys(popup_keys):
            send_combo(controller, part)
        # Screenshots only until the popup is there and the same twice in a
        # row (drawn), then its text is read once.
        box, previous, after_layout = None, None, []
        deadline = time.time() + 4
        while time.time() < deadline:
            time.sleep(0.15)
            shot = self._screenshot(hwnd)
            found = popup_box(image, shot)
            if found and previous is not None and not screens_differ(previous, shot):
                box = found
                after_layout = self._describe(shot)[4]
                break
            previous = shot if found else None
        lines = []
        if box:
            left, top, right, bottom = box
            inside = [line for line in after_layout if line.get("text", "").strip()
                      and left - 6 <= line["x"] + line.get("w", 0) / 2 <= right + 6 and top - 6 <= line["y"] <= bottom + 6]
            lines = [entry["text"] for entry in screen_rows(inside)]
        for part in list_loop_keys(close_keys):
            send_combo(controller, part)
        deadline = time.time() + 4
        while time.time() < deadline:
            time.sleep(0.15)
            if not popup_box(image, self._screenshot(hwnd)):
                break  # the list is back
        row_popups[key] = lines
        if not lines:
            self.say(f"{label}: {key}: no popup seen after {popup_keys}")

    def _write_collected_csv(self, collected: dict, value_label: str, label: str):
        """The collected rows with their value and popup info as a CSV
        ("lot;epp;info"), in the action's dump folder (else outputs/) -- the
        file for SnappySjaak's EPP list upload. The same columns are added to
        the run's dump file (Ctrl+P), matched by lot."""
        rows = [item for item in collected.get("items", []) if item.get("value") or item.get("info")]
        if not rows:
            self.say(f"{label}: no {value_label} / popup read on any row -- no CSV written.")
            return
        with_info = any(item.get("info") for item in rows)
        target_dir = self._export_dir()
        target = target_dir / f"{day_list_slug(collected['name'])}_{self.run_date():%Y%m%d}.csv"
        header = f"lot;{value_label.lower()}" + (";info" if with_info else "")
        lines = [header] + [
            f"{item['key']};{item.get('value', '')}" + (f";{str(item.get('info', '')).replace(';', ',')}" if with_info else "")
            for item in rows
        ]
        target.write_text("\r\n".join(lines) + "\r\n", encoding="utf-8", newline="")
        self.exported_files.append(str(target))
        self.collected_csv_paths.add(str(target))
        self.say(f"{label}: {len(rows)} row(s) -> {target}")
        if not self._merge_into_dump(rows, value_label, with_info, label):
            # Uniware's folder dump comes at the end of the run: then.
            self.pending_dump_merge = (rows, value_label, with_info, label)

    def _merge_into_dump(self, rows: list[dict], value_label: str, with_info: bool, label: str) -> bool:
        """Add the collected columns to this run's dump -- the Excel file from
        Ctrl+P, or Uniware's own CSV in its dump folder (dump.filipe....csv):
        the row whose cell is the lot gets them. False: no dump (yet)."""
        dumps = [Path(path) for path in self.exported_files
                 if str(path).lower().endswith((".xlsx", ".csv")) and str(path) not in self.collected_csv_paths]
        if not dumps:
            return False
        dump = dumps[-1]
        if dump.suffix.lower() == ".csv":
            self._merge_into_csv_dump(dump, rows, value_label, with_info, label)
            return True
        try:
            import openpyxl  # only needed here
            book = openpyxl.load_workbook(dump)
            sheet = book.active
            by_lot = {str(item["key"]): item for item in rows}
            first_free = sheet.max_column + 1
            sheet.cell(row=1, column=first_free, value=value_label)
            if with_info:
                sheet.cell(row=1, column=first_free + 1, value="Popup")
            filled = 0
            for row_cells in sheet.iter_rows(min_row=2):
                lot = next((str(cell.value).strip() for cell in row_cells if cell.value is not None and str(cell.value).strip() in by_lot), None)
                if not lot:
                    continue
                item = by_lot[lot]
                sheet.cell(row=row_cells[0].row, column=first_free, value=item.get("value", ""))
                if with_info:
                    sheet.cell(row=row_cells[0].row, column=first_free + 1, value=item.get("info", ""))
                filled += 1
            book.save(dump)
            _relative_sheet_paths(dump)
            self.say(f"{label}: {filled} row(s) of the dump {dump.name} got {value_label}" + (" and Popup" if with_info else ""))
        except Exception as error:  # noqa: BLE001 -- the CSV is there anyway
            self.say(f"{label}: could not add the columns to the dump {dump.name}: {error}")
        return True

    def _merge_into_csv_dump(self, dump: Path, rows: list[dict], value_label: str, with_info: bool, label: str):
        """Uniware's ";" CSV dump: after its header row (the one with "Prt")
        every row whose cell is a collected lot gets the value and Info."""
        try:
            raw = dump.read_bytes()
            encoding = "utf-8-sig" if raw.startswith(b"\xef\xbb\xbf") else "cp1252"
            table = list(csv.reader(raw.decode(encoding, errors="replace").splitlines(), delimiter=";"))
            header_index = next((index for index, cells in enumerate(table[:40]) if any(cell.strip().lower() == "prt" for cell in cells)), None)
            if header_index is None:
                self.say(f"{label}: the dump {dump.name} has no Prt column -- the columns are only in the CSV next to it.")
                return
            by_lot = {str(item["key"]): item for item in rows}
            prt = next(index for index, cell in enumerate(table[header_index]) if cell.strip().lower() == "prt")
            table[header_index] = table[header_index] + [value_label] + (["Popup"] if with_info else [])
            filled = 0
            # The dump has every row: the lots the loop didn't read.
            dump_lots = [cells[prt].strip() for cells in table[header_index + 1:] if prt < len(cells) and re.fullmatch(r"\d{6,10}", cells[prt].strip())]
            missing = [lot for lot in dump_lots if lot not in by_lot]
            self.say(f"{label}: the dump has {len(dump_lots)} lot(s), the loop read {len(dump_lots) - len(missing)}"
                     + (f" -- not read: {', '.join(missing[:20])}{' ...' if len(missing) > 20 else ''}" if missing else " -- all of them"))
            for cells in table[header_index + 1:]:
                lot = next((cell.strip() for cell in cells if cell.strip() in by_lot), None)
                if not lot:
                    continue
                item = by_lot[lot]
                cells.extend([item.get("value", "")] + ([item.get("info", "")] if with_info else []))
                filled += 1
            temp = dump.with_suffix(".tmp")
            with open(temp, "w", encoding=encoding, newline="") as handle:
                csv.writer(handle, delimiter=";", lineterminator="\r\n").writerows(table)
            temp.replace(dump)
            self.say(f"{label}: {filled} row(s) of the dump {dump.name} got {value_label}" + (" and Popup" if with_info else ""))
        except Exception as error:  # noqa: BLE001 -- the CSV is there anyway
            self.say(f"{label}: could not add the columns to the dump {dump.name}: {error}")

    def _press(self, controller, key: str, times: int):
        for _ in range(times):
            send_combo(controller, key)
            time.sleep(0.05)

    def _choose_row(self, hwnd, controller, target: str, label: str):
        """Move the highlight in a list to the row showing `target`: read the
        list, count the rows between the highlight and that row, press
        Down/Up that often, read again to check. Rows not on screen: scroll
        down to the end of the list, then up."""
        if not target.strip():
            raise RuntimeError(f"{label}: no value to look for in the list (fill in the field).")
        direction, previous_text = "Down", None
        for _ in range(60):
            self._check(hwnd)
            image, _, _, _, text, layout = self._read(hwnd, timeout=1.5)
            rows = list_items(layout, target)
            highlighted = [index for index, row in enumerate(rows) if row["highlighted"]]
            if not highlighted:
                self._save_problem_screen(image)
                raise RuntimeError(f"{label}: can't see which row is selected in the list. Stopped (screenshot in outputs).")
            current = highlighted[0]
            found = find_row(rows, target)
            if found == current:
                self.say(f"{label}: '{target}' selected ({rows[current]['text'][:60]})")
                return
            if found is not None:
                distance = found - current
                self._press(controller, "Down" if distance > 0 else "Up", abs(distance))
                previous_text = None
                continue
            # Not on screen: move past the visible rows to scroll the list.
            if text == previous_text:
                if direction == "Up":
                    break
                direction = "Up"  # end of the list: search the part above
            previous_text = text
            # Past the last (first) visible row and about a screen further.
            to_edge = len(rows) - 1 - current if direction == "Down" else current
            self._press(controller, direction, to_edge + max(1, len(rows) - 1))
        self._save_problem_screen(screen_reader.capture_window(hwnd))
        raise RuntimeError(f"{label}: '{target}' is not in the list today. Stopped (screenshot in outputs).")

    def _save_problem_screen(self, image):
        OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)
        name = f"{self.action['slug']}_{self.started:%Y%m%d_%H%M%S}_stopped.png"
        image.save(OUTPUTS_DIR / name)
        self.output_files.append(name)

    def _capture(self, hwnd, step_number: int, all_pages: bool, controller, reading):
        screen_id, title, _, text, layout = reading
        pages = [{"screen": screen_id, "title": title, "lines": layout}]
        previous = text
        while all_pages and len(pages) < 200:
            self._check(hwnd)
            send_combo(controller, "Page Down")
            # Wait until the next page is drawn (and stays the same); no
            # change at all within PAGE_WAIT_SECONDS = that was the last page.
            deadline = time.time() + PAGE_WAIT_SECONDS
            new_page = None
            while time.time() < deadline:
                self._check(hwnd)
                _, screen_id, title, _, text, layout = self._read(hwnd, timeout=1.5)
                if text != previous:
                    _, _, _, _, again, again_layout = self._read(hwnd, timeout=1.5)
                    if again == text:
                        new_page = (screen_id, title, text, again_layout)
                        break
                time.sleep(0.3)
            if new_page is None:
                break
            screen_id, title, previous, layout = new_page
            pages.append({"screen": screen_id, "title": title, "lines": layout})
        self.captures.append({"step": step_number, "pages": pages})
        self.say(f"Captured {len(pages)} page(s) of {title}")

    def run(self):
        self.active = True
        try:
            hwnd = find_uniware_window(self.title_part)
            if not hwnd:
                raise RuntimeError("Uniware window not found -- start Uniware (logged in) first.")
            if not bring_to_front(hwnd):
                raise RuntimeError("Could not bring Uniware to the front -- click it once and start again.")
            controller = keyboard.Controller()
            _, screen_id, title, selected, text, layout = self._read(hwnd)
            steps = self.action.get("steps", [])
            if steps and steps[0].get("menu"):
                # Starts with a menu route: any menu position is fine (the
                # route steers from there); from inside a program, go back
                # to the menu first.
                screen_id, title, selected, text, layout = self._back_to_menu(hwnd, controller, (screen_id, title, selected, text, layout))
            elif self.action.get("start_screen") and screen_id != self.action["start_screen"]:
                raise RuntimeError(f"Uniware is on '{title}', but this action starts on '{self.action.get('start_title') or self.action['start_screen']}'. Go there first.")
            self.say(f"Start: {title}")
            opens_window = any(step.get("opens_window") for step in steps)
            dump_before = None if opens_window else self._dump_files_now()
            if self.action.get("capture_start"):
                self._capture(hwnd, 0, False, controller, (screen_id, title, selected, text, layout))
            steps = self.action.get("steps", [])
            selected = self._run_steps(hwnd, controller, steps, 0, len(steps), selected)
            if dump_before is not None:
                self._collect_uniware_dump(dump_before)
                if self.pending_dump_merge:
                    self._merge_into_dump(*self.pending_dump_merge)
                    self.pending_dump_merge = None
            if self.dump_parts:
                self._join_dump_parts()
            self.result = "done"
            self.say("Finished.")
        except Exception as error:  # noqa: BLE001 -- shown in the dashboard
            self.result = "failed"
            self.say(f"STOPPED: {error}")
        finally:
            try:
                self._write_outputs()
            except OSError as error:
                self.result = "failed"
                self.say(f"Could not write the output: {error}")
            self.active = False

    def _run_steps(self, hwnd, controller, steps: list, start: int, end: int, selected):
        """Steps start..end-1; a "for each" step repeats the steps after it
        once per day-list item."""
        index = start
        while index < end:
            step = steps[index]
            number = index + 1
            self._check(hwnd)
            if step.get("foreach"):
                settings = step["foreach"]
                block_end = min(end, index + 1 + int(settings.get("count") or 1))
                self._run_foreach(hwnd, controller, steps, index + 1, block_end, settings, f"Step {number}")
                selected = self._read(hwnd)[3]
                index = block_end
                continue
            selected = self._run_one_step(hwnd, controller, step, number, selected)
            index += 1
        return selected

    def _run_foreach(self, hwnd, controller, steps, start, end, settings, label):
        data = load_day_list(settings["list"])
        items = [item for item in data.get("items", []) if not settings.get("only_open", True) or not item.get("label")]
        self.say(f"{label}: for each of {len(items)} order(s) in '{settings['list']}' -- steps {start + 1}-{end}")
        for position, item in enumerate(items, start=1):
            self.values[settings["field"]] = item["key"]
            self.current_item = (settings["list"], item["key"])
            self.say(f"{label}: ({position}/{len(items)}) {settings['field']} = {item['key']}")
            try:
                self._run_steps(hwnd, controller, steps, start, end, None)
            except Exception as error:
                data = load_day_list(settings["list"])
                entry = day_list_item(data, item["key"])
                if entry is not None:
                    entry["note"] = f"stopped: {error}"[:300]
                    save_day_list(data)
                raise
        self.current_item = None

    def _classify(self, hwnd, settings: dict, label: str):
        """The screen of the current day-list order (e.g. its product lines
        after F11): the word on a product line -> label_if, product lines
        without it -> label_else, both -> a warning (left for a person)."""
        item = getattr(self, "current_item", None)
        list_name = settings.get("list") or (item[0] if item else "")
        if not item:
            raise RuntimeError(f"{label}: a check step only works inside a 'for each order' block.")
        word = settings["word"]
        _, _, _, _, _, layout = self._read_stable(hwnd)
        product_rows = [row["text"] for row in screen_rows(layout) if PRODUCT_LINE_RE.search(row["text"])]
        with_word = [row for row in product_rows if re.search(re.escape(word), row, re.IGNORECASE)]
        without_word = [row for row in product_rows if row not in with_word]
        if with_word and without_word:
            result, note = "", f"order {item[1]} has multiple types of accessoires ({len(with_word)} {word}, {len(without_word)} other) -- check by hand"
        elif with_word:
            result, note = settings["label_if"], f"{word} ({len(with_word)} line(s))"
        elif without_word:
            result, note = settings["label_else"], f"no {word} ({len(without_word)} line(s))"
        else:
            result, note = "", "no product lines readable on this screen -- check by hand"
        data = load_day_list(list_name)
        entry = day_list_item(data, item[1])
        if entry is None:
            entry = {"key": item[1], "row": "", "type": ""}
            data.setdefault("items", []).append(entry)
        entry.update({"label": result, "note": note, "warning": not result, "checked": datetime.datetime.now().isoformat(timespec="seconds")})
        save_day_list(data)
        self.say(f"{label}: {item[1]} -> {result or 'WARNING'} ({note})")

    def _field_text(self, name: str, step: dict) -> str:
        """A field's typed text. {today} / {tomorrow} / {next_workday}, or a
        field filled in with one of those words (e.g. {date} = "tomorrow"),
        becomes that date -- written like the date typed while recording."""
        value = self.values.get(name, "")
        recorded = step.get("recorded_text") or ""
        day = date_for_word(name) or date_for_word(value)
        # A date filled in as 14-10-2026 / 2026-10-14 / 141026 ...: written
        # like the recorded one -- only on a step that typed a date when it
        # was recorded (never turns an order number into a date).
        if not day and value and re.search(r"(?<!\d)\d{2}[-/.]?\d{2}[-/.]?\d{2,4}(?!\d)", recorded):
            day = parse_date_value(value)
        if day:
            return format_date_like(day, recorded)
        # Left empty when starting: what was typed while recording (a field
        # that types nothing at all made Uniware stop at an empty input).
        return value if value != "" else str(step.get("recorded_text") or "")

    def _wait_not_busy(self, hwnd, label: str, max_wait: float = 300):
        """Uniware shows "Gegevens verzamelen ... Moment geduld aub..." while
        it loads, on the same screen (the screen name doesn't change): no key
        until that is gone."""
        started = time.time()
        next_note = started + 5
        while True:
            self._check(hwnd)
            _, _, title, _, text, _ = self._read(hwnd, timeout=1.5)
            if not BUSY_RE.search("\n".join([title, *text])):
                return
            waited = time.time() - started
            if waited > max_wait:
                raise RuntimeError(f"{label}: Uniware still says 'Moment geduld' after {max_wait:.0f}s. Stopped.")
            if time.time() > next_note:
                self.say(f"{label}: Uniware is loading (Moment geduld) -- waiting ({waited:.0f}s)")
                next_note = time.time() + 10
            time.sleep(0.5)

    def _run_one_step(self, hwnd, controller, step: dict, number: int, selected):
        if not step.get("foreach") and not step.get("classify"):
            self._wait_not_busy(hwnd, f"Step {number}")
        if step.get("loop"):
            self._run_list_loop(hwnd, controller, step["loop"], f"Step {number}")
            return self._read(hwnd)[3]
        if step.get("classify"):
            self._classify(hwnd, step["classify"], f"Step {number}")
            return selected
        if "text" in step:
            typed = re.sub(r"\{([a-z0-9_]+)\}", lambda match: self._field_text(match.group(1), step), step["text"])
            controller.type(typed)
            self.say(f"Step {number}: typed {typed!r}")
        else:
            if step.get("menu"):
                self._go_menu(hwnd, controller, step["menu"], f"Step {number}")
            elif step.get("pick"):
                target = re.sub(r"\{([a-z0-9_]+)\}", lambda match: self.values.get(match.group(1), ""), step["pick"])
                self._choose_row(hwnd, controller, target, f"Step {number}")
            elif step.get("enter_on"):
                # Arrow to the recorded menu item first if the
                # highlight is somewhere else.
                for _ in range(30):
                    if selected and selected[0] == step["enter_on"]:
                        break
                    send_combo(controller, "Down")
                    _, _, _, selected, _, _ = self._read(hwnd, timeout=1.5)
                else:
                    raise RuntimeError(f"Menu item '{step['enter_on']}' not found.")
            windows_before = visible_windows() if step.get("opens_window") else None
            send_combo(controller, step["key"])
            if windows_before is not None:
                self._take_window_dump(hwnd, windows_before, step, f"Step {number}")
                if not bring_to_front(hwnd):
                    raise RuntimeError("Could not bring Uniware back to the front after the dump window.")
            self.say(f"Step {number}: {step['key']}" + (f" on '{step['pick']}'" if step.get("pick") else f" on '{step['enter_on']}'" if step.get("enter_on") else ""))
        reading = self._wait_for(hwnd, step, f"Step {number}")
        screen_id, title, selected, text, layout = reading
        if step.get("read_column") and step.get("read_field"):
            self._read_column_value(hwnd, step, f"Step {number}", layout)
        if step.get("capture"):
            self._capture(hwnd, number, bool(step.get("capture_pages")), controller, reading)
        return selected

    def run_date(self) -> datetime.date:
        """The day this run is about: the {date} field when it was filled in
        (14-10-2026, yesterday ...), else the day it runs -- so a backfill of
        older days names its files after those days."""
        return parse_date_value(self.values.get("date", "")) or self.started.date()

    def export_name(self) -> str:
        """The dump file name from the action's template: {date} (the run's
        day, see run_date), {time}, {action} and any field, e.g.
        "tour-508_{date}"."""
        template = self.action.get("export_name") or "{action}_{date}_{time}"
        names = {**self.values, "action": self.action["slug"], "date": f"{self.run_date():%Y%m%d}", "time": f"{self.started:%H%M%S}"}
        name = re.sub(r"\{([a-z0-9_]+)\}", lambda match: names.get(match.group(1), ""), template)
        return re.sub(r'[\\/:*?"<>|]+', "-", name).strip(" .") or self.action["slug"]

    # Uniware writes its own dump file (e.g. "dump.filipe.2026-10-14_16_57_30.csv")
    # into a folder; the run picks up the new one and gives it the action's name.
    def _dump_files_now(self) -> dict | None:
        folder = (self.action.get("uniware_dump_dir") or "").strip()
        if not folder:
            return None
        if not Path(folder).is_dir():
            raise RuntimeError(f"Uniware's dump folder '{folder}' doesn't exist (see Edit).")
        pattern = self.action.get("uniware_dump_pattern") or "dump.*.csv"
        return {path.name: path.stat().st_mtime for path in Path(folder).glob(pattern) if path.is_file()}

    def _collect_uniware_dump(self, before: dict):
        folder = Path(self.action["uniware_dump_dir"].strip())
        pattern = self.action.get("uniware_dump_pattern") or "dump.*.csv"
        max_wait = float(self.action.get("uniware_dump_wait") or 180)
        started = time.time()
        next_note = started + 10
        last_size, steady_since = None, None
        self.say(f"Waiting for Uniware's dump file ({pattern}) in {folder}")
        while True:
            if self.stop_requested:
                raise RuntimeError("Stopped by you.")
            new = [path for path in folder.glob(pattern)
                   if path.is_file() and (path.name not in before or path.stat().st_mtime > before[path.name])]
            if new:
                newest = max(new, key=lambda path: path.stat().st_mtime)
                size = newest.stat().st_size
                if size > 0 and size == last_size:
                    # Same size for 3 s: Uniware (and OneDrive) are done
                    # writing. Moving fails while it's still open -- retry.
                    if time.time() - steady_since >= 3 and self._move_dump(newest):
                        return
                else:
                    last_size, steady_since = size, time.time()
            waited = time.time() - started
            if waited > max_wait:
                raise RuntimeError(f"No new dump file ({pattern}) appeared in {folder} within {max_wait:.0f}s.")
            if time.time() > next_note:
                self.say(f"Still waiting for the dump file ({waited:.0f}s of max {max_wait:.0f}s)")
                next_note = time.time() + 15
            time.sleep(1)

    def _move_dump(self, source: Path) -> bool:
        target_dir = Path((self.action.get("export_dir") or "").strip() or source.parent)
        target_dir.mkdir(parents=True, exist_ok=True)
        target = target_dir / f"{self.export_name()}{source.suffix}"
        try:
            if target_dir.resolve().drive == source.resolve().drive:
                source.replace(target)
            else:
                temp = target_dir / f"~{target.name}.tmp"
                temp.write_bytes(source.read_bytes())
                temp.replace(target)
                source.unlink()
        except PermissionError:
            return False  # still being written
        self.exported_files.append(str(target))
        self.say(f"Dump file {source.name} -> {target}")
        return True

    def _write_outputs(self):
        if not self.captures:
            return
        OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)
        stamp = f"{self.action['slug']}_{self.started:%Y%m%d_%H%M%S}"
        text_path = OUTPUTS_DIR / f"{stamp}.txt"
        csv_path = OUTPUTS_DIR / f"{stamp}.csv"
        with text_path.open("w", encoding="utf-8") as handle:
            for capture in self.captures:
                for index, page in enumerate(capture["pages"], start=1):
                    handle.write(f"=== step {capture['step']} -- {page['title']} -- page {index}\n")
                    for line in page["lines"]:
                        handle.write(f"{line['text']}\n")
        with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle, delimiter=";")
            writer.writerow(["step", "page", "screen", "y", "x", "text"])
            for capture in self.captures:
                for index, page in enumerate(capture["pages"], start=1):
                    for line in page["lines"]:
                        writer.writerow([capture["step"], index, page["screen"], line["y"], line["x"], line["text"]])
        (OUTPUTS_DIR / f"{stamp}.json").write_text(json.dumps(self.captures, indent=2, ensure_ascii=False), encoding="utf-8")
        self.output_files = [text_path.name, csv_path.name] + self.output_files
        export_dir = (self.action.get("export_dir") or "").strip()
        if not export_dir or self.action.get("uniware_dump_dir"):
            return  # with Uniware's own dump file, that file is the dump
        if self.result != "done":
            self.say("NOT put in the dump folder: the run did not finish (the partial data is only in outputs).")
            return
        target_dir = Path(export_dir)
        target_dir.mkdir(parents=True, exist_ok=True)
        base = self.export_name()
        for source, suffix in ((csv_path, ".csv"), (text_path, ".txt")):
            target = target_dir / f"{base}{suffix}"
            # Written under a temporary name first, so OneDrive never syncs a
            # half-written file.
            temp = target_dir / f"~{base}{suffix}.tmp"
            temp.write_bytes(source.read_bytes())
            temp.replace(target)
            self.exported_files.append(str(target))
        self.say(f"Dump written to {target_dir}: {base}.csv / .txt")
